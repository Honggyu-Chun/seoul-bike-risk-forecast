"""Generate the plain-language handout with rationale, explicit metrics and rubric mapping."""
from pathlib import Path
from datetime import datetime, timezone
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables'
OUT.mkdir(exist_ok=True)
FONT = 'Noto Sans CJK KR'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.65), Cm(1.55)
sec.left_margin, sec.right_margin = Cm(1.8), Cm(1.8)
sec.header_distance, sec.footer_distance = Cm(.65), Cm(.7)

def run(p, text, size=10.5, bold=False, color='222222'):
    r = p.add_run(text)
    r.font.name = FONT
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    r.font.size = Pt(size)
    r.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    return r

for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    s = doc.styles[name]
    s.font.name = FONT
    s.font.size = Pt(10.5)
    s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    s.paragraph_format.line_spacing = 1.18
    s.paragraph_format.space_after = Pt(5)
    s.font.color.rgb = RGBColor.from_string('222222')
header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run(header, 'AI P-실무프로젝트  |  주제 선정 회의자료', 8, color='666666')
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(footer, '- ', 8, color='666666')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
run(footer, ' -', 8, color='666666')

def para(text, size=10.5, bold=False, color='222222', after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    run(p, text, size, bold, color)
    return p

def bullet(label, text, sub=False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(.65 if sub else .15)
    p.paragraph_format.first_line_indent = Cm(-.3 if sub else -.15)
    p.paragraph_format.space_after = Pt(5)
    run(p, '- ' if sub else '○ ')
    if label:
        run(p, label + ' ', bold=True)
    run(p, text)
    return p

def note(text): return para('※ ' + text, 8.5, color='666666', after=4)

def section(num, title, new_page=False):
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = new_page
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(11)
    pp = p._p.get_or_add_pPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'DEDEDE'); pp.append(sh)
    borders = OxmlElement('w:pBdr')
    for side in ['top','bottom']:
        el = OxmlElement('w:' + side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'5'); el.set(qn('w:color'),'777777'); el.set(qn('w:space'),'6'); borders.append(el)
    pp.append(borders)
    run(p, f'  {num}.  {title}', 13, True)

def sub(title):
    p = para('□ ' + title, 11, True, after=6)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.keep_with_next = True

def table(head, rows, widths, size=9.3):
    t = doc.add_table(rows=1, cols=len(head))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col,w in zip(t.columns,widths): col.width = Cm(w)
    for row_i, values in enumerate([head]+rows):
        cells = t.rows[0].cells if row_i == 0 else t.add_row().cells
        for j, text in enumerate(values):
            c=cells[j]; c.width=Cm(widths[j]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0]
            p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(4)
            p.paragraph_format.line_spacing=1.12
            run(p,text,size,row_i==0)
            if row_i==0: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            cp=c._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E5E5E5' if row_i==0 else ('F5F5F5' if row_i%2 else 'FFFFFF')); cp.append(sh)
            bd=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'C5C5C5'); bd.append(e)
            cp.append(bd)
        cant=OxmlElement('w:cantSplit'); t.rows[row_i]._tr.get_or_add_trPr().append(cant)
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    return t

def box(title,text):
    t=table([title],[(text,)],[17.4],10)
    t.rows[0].cells[0].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.LEFT
    return t

def hyperlink(p,label,url):
    rel=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rel)
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    for tag,attr,value in [('color','val','444444'),('sz','val','16'),('u','val','single')]:
        el=OxmlElement('w:'+tag); el.set(qn('w:'+attr),value); pr.append(el)
    rf=OxmlElement('w:rFonts'); rf.set(qn('w:ascii'),FONT); rf.set(qn('w:eastAsia'),FONT); pr.append(rf)
    r.append(pr); tx=OxmlElement('w:t'); tx.text=label; r.append(tx); h.append(r); p._p.append(h)

p=para('따릉이 이용량 예측 서비스 검토(안)',18,True,after=9)
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=para('2026. 10. 19.  /  AI P-실무프로젝트 주제 검토',9.5,color='666666',after=10)
p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
box('어디에 쓰는가', '따릉이 운영 담당자가 출퇴근 시간대의 자전거 배치를 검토할 때, 이용이 몰릴 것으로 예상되는 대여소를 먼저 추리는 데 사용함. 서비스는 예상 대여 건수가 많은 곳의 목록과 근거를 제공하고, 담당자는 현재 자전거 수·반납 상황 등을 함께 보고 실제 조치를 결정함. 첫 결과물은 과거 날짜로 이 과정을 시연하는 화면임.')

section('1','왜 해야 하며, 어떤 업무에 쓰는가')
sub('가. 지금 남은 자전거 수와 앞으로 빌려갈 건수는 다른 정보임')
bullet('현재 확인 가능한 정보', '서울시 공개 서비스는 대여소 위치와 지금 빌릴 수 있는 자전거 수 등을 제공함.[1]')
bullet('해결하려는 문제', '현재 자전거 수가 같아도 앞으로의 이용 규모는 다를 수 있음. 정기 운영계획을 세울 때 현재 수량만 보면, 어느 곳의 이용이 늘어날지 비교하기 어려움.')
box('이해를 돕기 위한 가상 사례', '현재 자전거가 10대씩 있는 두 대여소가 있다고 가정함. 한 곳은 다음 한 시간에 2건, 다른 곳은 15건의 대여가 예상된다면 담당자가 관심 있게 볼 곳은 달라질 수 있음. 단, 실제 부족 여부는 반납과 자전거 이동까지 함께 확인해야 함.')
sub('나. 활용처는 “자전거 배치 검토 대상 선정”으로 구체화함')
bullet('누가·언제', '운영 담당자가 정기 운영계획에서 요일별 출퇴근 시간대의 배치·점검 대상을 살펴볼 때 사용함.')
bullet('무엇을 결정하는가', '서울 전역을 일일이 살펴보기 전에, 예상 이용량이 많은 대여소 목록을 받아 먼저 확인할 곳을 정함. 이동할 자전거 수까지 자동으로 결정하는 기능은 아님.')
bullet('왜 필요한가', '미리 살펴볼 곳을 추리는 것이 목적이므로, 건수 오차뿐 아니라 실제 이용이 많았던 곳을 잘 골랐는지도 평가함. 실제 담당자의 업무에 맞는지는 의견을 받아 확인해야 함.')
sub('다. 이번 과제에서는 “대여 건수 예측”부터 시작함')
bullet('왜 AI를 시험하는가', '요일·시간·날씨와 주변 이용 흐름을 함께 보면 과거 평균보다 잘 맞을 가능성이 있음. 실제로 평균만큼도 못 맞히면 복잡한 AI를 쓸 근거가 약해지므로, 기본 방법과 비교하는 것이 과제의 핵심임.')
note('이 문서의 “이용량”은 실제 대여가 이루어진 건수임. 이용이 많다는 것만으로 자전거 부족을 뜻하지 않으며, 실제 운영 효과는 아직 검증하지 않았음.')

section('2','기존 방식과 무엇이 다른가',True)
sub('가. 현재 수량 조회에 더해, 시간대별 예상 이용량을 보여줌')
table(['구분','주로 답하는 질문','확인하거나 만들 내용'],[
    ('서울시 공개 대여정보 [1]','“지금 어디에서 몇 대를 빌릴 수 있는가?”','현재 대여가능 수량과 위치를 확인하는 정보'),
    ('이번에 제안하는 서비스','“출퇴근 시간에 먼저 확인할 대여소는 어디인가?”','예상 이용량이 많은 대여소 목록·지도와 선정 근거를 제공. 실제 부족 여부는 담당자가 다른 정보와 함께 확인'),
],[3.7,6.7,7.0],9.7)
note('공식 공개 안내에서 확인한 기능과 비교한 것임. 운영기관 내부에서 사용하는 예측 방식까지 조사한 것은 아니므로, 기존에 예측 기능이 전혀 없다고 단정하지 않음.')
sub('나. 해당 대여소의 기록뿐 아니라 주변 이용 흐름도 참고함')
bullet('기본적인 계산', 'A대여소가 예전에 같은 요일·시간에 얼마나 이용되었는지, 그때의 날씨와 휴일 여부는 어땠는지 참고하여 예상 건수를 계산함.')
bullet('추가로 시험할 계산', 'A대여소 주변이나 이용 흐름이 비슷한 다른 대여소의 과거 기록도 함께 참고함. “이곳만 볼 때보다 주변도 함께 보면 더 잘 맞는가”를 확인하는 것임.')
bullet('이렇게 생각한 이유', '같은 지하철역이나 업무지역 주변의 대여소는 비슷한 시간에 이용이 늘어날 수 있음. 실제로 도움이 되는지는 같은 조건의 비교 시험으로 확인해야 함.')
bullet('주장할 차이', '주변 기록을 더하면 건수뿐 아니라 먼저 확인할 대여소도 더 잘 고르는지 검증함. 새로운 AI 발명이나 재배치 효과로 과장하지 않음.')
sub('다. 담당자는 다음 순서로 사용하게 됨')
table(['순서','담당자가 하는 일과 얻는 정보'],[
    ('① 요일·시간 선택','예상 대여 건수가 많은 대여소 목록을 지도와 함께 확인함. 비교 시험에서는 상위 50곳을 기준으로 함.'),
    ('② 확인 대상 검토','목록에서 대여소를 누르면 예상 건수·이용이 몰릴 시간·과거 예측 오차를 확인함.'),
    ('③ 실제 운영 판단','담당자가 기존 수량 정보와 반납·현장 상황을 추가 확인한 뒤 배치·점검 여부를 결정함. 첫 화면에는 검토 목록과 근거를 제공함.'),
    ('④ 지난 선정 결과 확인','실제로 이용이 많았던 대여소와 선정 목록을 비교하여, 먼저 볼 곳을 잘 골랐는지 확인함.'),
],[4.0,13.4],9.7)
note('위 화면은 앞으로 만들 기능의 구상임. 서울시 대여 건수 기록은 반년 단위로 갱신되므로 첫 시연은 과거 날짜의 예측과 실제 결과를 비교하는 방식으로 시작함. 현재 상황을 즉시 반영하려면 최신 이용 기록을 별도로 받아야 함.')

section('3','실제로 쓸 자료가 있는가',True)
sub('가. 핵심 자료 3종을 실제로 내려받아 확인함')
table(['자료·출처','무엇을 확인했는가','어디에 쓰는가'],[
    ('따릉이 이용 기록\n서울시 [2]','2025년 1월 기록을 확인함.\n대여소 번호 2,730개가 등장하며,\n실제 대여 건수 합계는 약 170만 건임.','어느 대여소에서 몇 시에 몇 건 대여되었는지 확인하고 예측을 연습·평가하는 데 사용'),
    ('대여소 위치\n서울시 [3]','2024년 12월 위치 자료와 대조함.\n2,730개 중 2,728개의 위치를 찾음.\n이 중 2,674곳은 500m 안에 이웃 대여소가 있음.','지도에 대여소를 표시하고, 가까운 대여소끼리 연결하는 데 사용'),
    ('기온·습도·바람 등\n기상청 [4]','서울 대표 관측소 1곳의 2025년 자료를 확보함. 1월의 모든 시간대에 해당하는 날씨가 있고, 기온·습도·풍속 값은 빠짐없이 있음.','날씨에 따라 이용량이 어떻게 달라지는지 확인. 관측소 1곳으로 서울을 대표하는 한계가 있음'),
],[3.7,7.8,5.9],9.5)
sub('나. 사이트에 있다는 사실만 확인한 것이 아니라 내용을 대조함')
bullet('파일 열기', '공식 기관의 파일을 내려받아 읽을 수 있는지, 필요한 날짜·시간·대여소 번호·이용 건수가 있는지 확인함.')
bullet('건수 계산', '이용 기록은 나이·이용권 종류 등에 따라 여러 줄로 나뉘어 있음. 줄 수를 사람 수로 세지 않고, 파일에 적힌 “이용건수”를 더함. 합계는 1,704,190건임.')
bullet('위치·날씨 연결', '대여소 번호로 위치를 찾고, 같은 날짜·시간의 이용 기록과 날씨를 묶음. 1월 31일×24시간인 744개 시간대에 날씨가 모두 연결됨.')
note('확인 결과는 자료를 함께 쓸 수 있다는 뜻임. 모든 대여소가 매시간 빠짐없이 기록되었다거나, AI가 이미 정확해졌다는 뜻은 아님. 대여소 수는 당시 자료 기준임.')
sub('다. 추가로 찾아본 자료와 남은 준비')
bullet('지하철·공원', '공식 자료를 실제로 열어 확인함.[6][7][8] 역별 승하차 인원이나 공원과의 거리를 추가할 수 있으나, 첫 단계의 필수 자료로 정하지는 않음.')
bullet('공휴일·계절', '최신 2025년 공식 공휴일 목록은 추가 확보가 필요함.[9] 또 겨울 한 달로 1년 전체를 평가할 수 없으므로 다른 계절의 이용 기록을 더 확보해야 함.')
note('해외 대학의 서울 따릉이 공개 자료도 확인했으나 대여소 번호가 없어, 대여소별 지도를 만드는 주자료로는 사용하지 않음.[5]')

section('4','정확히 어떤 평가지표를 쓸 것인가',True)
sub('가. “몇 건인가”를 맞히는 문제이므로 숫자의 차이로 평가함')
bullet('문제의 종류', '“이용이 많다/적다”를 고르는 문제가 아니라 “한 시간에 몇 건 대여되는가”를 계산하는 문제임. 이런 숫자 예측을 회귀라고 하며, 예상 건수와 실제 건수의 차이로 평가함.')
table(['지표','쉬운 뜻과 선택 이유','이번 과제의 용도'],[
    ('MAE\n평균 절대오차','각 예측에서 몇 건 틀렸는지 더해 평균을 냄. “평균 3건 틀림”처럼 설명함.','주평가지표. 작을수록 좋음. 기본 AI보다 건수 예측이 나아졌는지 판단'),
    ('RMSE\n평균제곱근오차','큰 실수에 더 큰 무게를 둠. 일부 대여소에서 크게 틀리면 값이 커지며, 건수 단위로 읽을 수 있음.','보조평가지표. 작을수록 좋음. 큰 실수가 늘지 않았는지 점검'),
    ('상위 50곳 적중률\nPrecision@50','예상 이용량 상위 50곳 중 실제 이용량 상위 50곳에 든 비율. 50곳 중 30곳이 겹치면 60%임.','업무 용도 확인 지표. 클수록 좋음. 먼저 살펴볼 대여소를 잘 골랐는지 판단'),
    ('L2 계열: MSE\n평균제곱오차','틀린 건수를 제곱해 평균낸 값. 큰 실수를 줄이도록 AI를 학습시킬 때 사용. RMSE는 이 값의 제곱근임.','학습 기준. 발표 점수는 MAE·RMSE로 설명'),
    ('F1 점수','“이용 많음/보통” 같은 분류의 적중과 놓침을 함께 평가. 몇 건 틀렸는지는 알 수 없음.','이번 주지표로 미채택. 건수는 MAE·RMSE, 목록 선정은 상위 50곳 적중률로 평가'),
],[3.4,8.5,5.5],9.2)
sub('나. 계산 예시: 실제 성능이 아닌 설명용 숫자임')
para('실제 대여가 각각 10건인 두 시간에 8건, 14건으로 예상했다면 틀린 크기는 2건과 4건임.\nMAE = (2 + 4) ÷ 2 = 3건\nMSE = (2² + 4²) ÷ 2 = 10건²   /   RMSE = √10 ≈ 3.16건',10,after=5)
note('평가 1건은 “한 대여소의 한 시간”임. 전체 MAE는 모든 평가 건의 오차 크기를 더해 평가 건수로 나눈 값, RMSE는 오차 제곱을 같은 방식으로 평균낸 뒤 제곱근을 구한 값임.')
note('오차는 양·음으로 상쇄하지 않음. 여기서 학습용 L2는 MSE 방식이며, 모델 복잡도를 제한하는 L2 정규화와는 구분함.')
sub('다. 0건도 있는 자료에 맞게 지표를 선택함')
bullet('지표 선택 원칙', '확인된 0건도 평가함. 실제값으로 나누는 MAPE는 0건에서 계산이 어려워 주지표로 쓰지 않음. 두 목록을 똑같이 50곳씩 고르는 이번 비교에서는 F1을 별도로 계산해도 상위 50곳 적중률과 값이 같음.')

section('5','비교 시험과 정량 목표를 어떻게 정할 것인가',True)
sub('가. 쉬운 방법과 충분히 강한 AI를 모두 비교 대상으로 둠')
table(['비교 방법','사용하는 자료','비교의 의미'],[
    ('① 과거 평균\n단순 기준','같은 대여소의 공개된 과거 동일 요일·시간 평균','AI가 간단한 평균보다 나은지 확인'),
    ('② 기본 AI\n주 비교 기준','해당 대여소 과거 기록 + 요일·휴일 + 이용 가능한 날씨. LightGBM이라는 공개 예측 프로그램 사용','제안 방법과 직접 비교할 기본 방법(Baseline)'),
    ('③ 제안 AI\n개선 후보','②와 같은 프로그램에 주변·이용 흐름이 비슷한 대여소의 과거 기록을 추가','프로그램을 바꾼 효과가 아니라 주변 기록을 추가한 효과를 확인'),
],[3.6,8.4,5.4],9.2)
sub('나. 회의용 정량 목표(안)')
box('팀의 목표(안)이며, 학교가 지정한 합격선은 아님', '제안 AI는 기본 AI보다 ①전체 MAE 5% 이상 감소 ②RMSE 악화 없음 ③출퇴근 시간 상위 50곳 적중률 5%p 이상 향상을 목표로 함. 단순 평균과도 비교하며, 학습·검증 결과로 가능성을 점검한 뒤 최종 시험 전에 목표를 확정함.')
para('MAE 개선율 = (기본 AI MAE − 제안 AI MAE) ÷ 기본 AI MAE × 100\n예: 4.0건 → 3.8건은 5% 개선. 적중률 60% → 65%는 5%p 향상임. 실제 성능이 아님.',9.5,after=5)
note('기본 AI의 MAE가 0이면 개선율은 계산하지 않고 두 방법의 오차를 그대로 보고함. 목표 미달 시에는 개선이 제한적이었다고 보고하며 시험 결과에 맞춰 목표를 낮추지 않음.')
sub('다. 유리한 자료만 골라 점수를 높이지 않도록 함')
bullet('같은 조건', '공개된 과거로 학습·조정하고 뒤의 기간으로 시험함. 같은 날짜·대여소·평가 건수와 AI 조정 조건을 적용함. 미래 날씨·기록이나 시험 자료로 만든 평균은 사용하지 않음.')
bullet('0건·누락 구분', '운영 여부와 누락을 확인해 평가 범위를 정함. 확인된 0건은 포함하고, 알 수 없는 시간은 별도 표시하며 포함·제외 수와 이유를 공개함.')
bullet('업무 용도 시험', '평일 7·8·18·19시마다 같은 대상에서 예상·실제 상위 50곳을 뽑아 겹친 수÷50을 구한 뒤 평균냄. 평가 대상은 50곳 이상이며 동률은 대여소 번호순으로 처리함.')
bullet('범위 고정', '50곳은 시험용 제안 규모임. 실제 업무에서 볼 수 있는 수는 담당자와 협의하되 최종 시험 전에 고정함. 전체·출퇴근·고이용 상위 20% 대여소의 오차와 주별 결과도 함께 공개함.')

section('6','공식 평가 항목과 결과물을 어떻게 연결할 것인가',True)
note('근거: 첨부 「2026_AI P-실무프로젝트_평가_260921.pdf」 2~3쪽 및 과제 계획서 양식 3·6항. 공식 문서는 특정 지표를 지정하지 않으며, 기본 방법 대비 우수성의 객관적 검증과 정량 목표를 요구함.')
sub('가. 사전평가 40%: 필요성과 실행 계획을 설명함')
table(['공식 평가 항목','이 과제에서 보여줄 근거'],[
    ('주제 적합성','정기 배치 계획 검토 때 먼저 살펴볼 대여소를 추린다는 활용처, 현재 수량과 미래 대여 건수의 차이, 실제 정답이 있는 AI 과제라는 점을 설명'),
    ('창의성','대여소별 예측에 주변 기록을 더했을 때 언제·어디서 나아지는지 확인한다는 구체적인 질문을 제시. 주변 정보 자체를 새로운 발명으로 과장하지 않음'),
    ('발표 구성·질의응답 / 계획서','실제 확보한 자료, MAE·RMSE 선택 이유, 비교 대상·시험 규칙·정량 목표, 역할과 일정을 계획서 및 발표에 동일하게 반영'),
],[5.0,12.4],9.1)
sub('나. 최종평가 60%: 점수·작동 화면·재현 근거를 제시함')
table(['공식 평가 항목','남길 결과물과 확인 기준'],[
    ('모델 성능','전체·출퇴근·고이용 대여소의 MAE·RMSE와 평가 건수, 크게 틀린 사례를 제시. 단순한 “정확도 몇 %” 대신 건수 오차로 설명'),
    ('벤치마크 대비 우수성\n(기본 방법과 객관적 비교)','①평균·②기본 AI·③제안 AI를 같은 기간으로 비교. MAE 개선율·RMSE·상위 50곳 적중률의 목표 달성 여부와 불리한 결과도 공개'),
    ('구현 완성도','지도 조회 → 대여소 그래프 → 과거 실제값 비교의 3개 사용 절차를 모두 끝까지 시연하는 것을 목표로 함(3개 중 3개 작동)'),
    ('적용 가능성','상위 50곳 적중률로 후보 선정의 질을 확인하고, 배치 검토 과정에 목록·근거가 유용한지 사용자 의견을 수집. 실제 담당자 참여 여부를 밝히고 재배치 효과로 확대 해석하지 않음'),
    ('문서화·오픈소스화','자료 출처·버전, 0건/누락 처리, 기간 분리, 지표 계산식, 실행 방법과 코드를 GitHub/Wiki에 공개. 같은 자료로 결과 재현 가능하도록 정리'),
    ('발표 구성·질의응답','왜 해야 하는지 → 자료 근거 → 비교 조건 → 결과 → 한계 순서로 설명. 가상 예와 실제 성능을 명확히 구분'),
    ('동료평가','자료·모델·화면·문서별 담당과 완료 결과를 기록하여 팀원 기여를 확인할 수 있도록 함'),
],[5.0,12.4],9.1)
note('항목별 추가 배점이나 MAE 5% 합격선이 공식 안내에 정해진 것은 아님. 위 내용은 공개된 평가 항목에 대응하기 위한 팀의 설계안이며, 점수를 보장하는 의미가 아님.')

section('7','오늘 회의에서 확정할 사항',True)
sub('가. 서비스·지표·담당을 함께 결정함')
bullet('활용처', '“운영 담당자의 출퇴근 시간대 배치 검토를 위한 대여소 후보 선정”으로 용도를 정할지 결정함. 실제 담당자의 업무·수요를 확인할 연락 또는 검토 경로도 정함.')
bullet('지표와 목표', '기본 AI 대비 MAE 5% 감소·RMSE 악화 없음·상위 50곳 적중률 5%p 향상의 제안 목표를 검토함. 학습은 L2 계열 MSE를 쓰고 F1은 별도로 쓰지 않음.')
bullet('첫 개발 범위', '지도·예상 건수 그래프·과거 실제값 비교의 3개 기능부터 만들지 결정함. 지하철·공원 정보와 현재 상황의 즉시 반영은 후속 단계로 둘 수 있음.')
bullet('담당', '① 자료 정리·확인 ② 예측 계산·정확도 비교 ③ 계산 결과를 화면에 전달 ④ 화면·발표·문서의 담당자와 점검일을 정함.')
para('회의 결정: __________________________________________________________\n담당 및 다음 점검일: __________________________________________________',9.5,after=5)
sub('나. 가까운 제출 일정')
para('10/19~21  주제·개발 범위·담당 확정, 추가로 필요한 자료 정리\n10/22~24  계획서·발표자료·화면 예시 작성 및 검토\n10/25(일) 23:59  사전평가 제출: PPT·PDF·10분 이내 영상의 주소 TXT·계획서 WORD를 압축하여 팀장 1명이 제출\n12/13(일) 23:59  최종평가 제출: 정확도 비교·서비스·결과보고서 등 정리',9,after=5)
note('공휴일 보완은 [9]의 “공휴일 정보 조회”에서 2025년 전체 목록을 받으면 됨. 나중에 지정된 1/27·6/3도 포함되었는지 확인하며, 인증키가 아닌 조회 결과 파일만 공유함.')
sub('참고 자료  ※ 제목 클릭 시 공식 출처로 이동')
refs=[
 ('[1] 서울시 실시간 대여정보','https://data.seoul.go.kr/dataList/OA-15493/A/1/datasetView.do'),
 ('[2] 시간대별 이용정보','https://data.seoul.go.kr/dataList/OA-15245/F/1/datasetView.do'),
 ('[3] 대여소 위치','https://data.seoul.go.kr/dataList/OA-13252/F/1/datasetView.do'),
 ('[4] 기상청 ASOS','https://data.kma.go.kr/data/grnd/selectAsosList.do?pgmNo=36'),
 ('[5] UCI Seoul Bike','https://archive.ics.uci.edu/dataset/560/seoul+bike+sharing+demand'),
 ('[6] 지하철 일별 승하차','https://data.seoul.go.kr/dataList/OA-12914/F/1/datasetView.do'),
 ('[7] 지하철 시간대별 월 누계','https://data.seoul.go.kr/dataList/OA-12252/S/1/datasetView.do'),
 ('[8] 서울 주요 공원','https://data.seoul.go.kr/dataList/OA-394/S/1/datasetView.do'),
 ('[9] 한국천문연구원 특일 정보','https://www.data.go.kr/data/15012690/openapi.do'),
]
for i in range(0,len(refs),3):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3)
    for j,(label,url) in enumerate(refs[i:i+3]):
        if j: run(p,'  |  ',8,color='666666')
        hyperlink(p,label,url)

props=doc.core_properties
props.title='따릉이 이용량 예측 서비스 검토(안) — 활용처·평가지표 보완본'
props.subject='배치 검토 대상 선정·MAE/RMSE/상위 50곳 적중률·L2/F1 선택·공식 평가 대응'
props.author='P_Project'
props.keywords='따릉이, 추진 검토, 회의자료, 데이터 검증, 수요 예측'
props.created=props.modified=datetime(2026,10,19,tzinfo=timezone.utc)
path=OUT/'따릉이_이용량예측_회의자료_활용처_평가지표보완_20261019.docx'
doc.save(path)
print(path)
