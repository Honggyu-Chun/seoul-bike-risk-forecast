"""Create a spacious, evidence-led rider service meeting handout."""
from datetime import datetime, timezone
import json
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import meeting_doc_format as base

doc, OUT, ROOT = base.doc, base.OUT, base.ROOT
audit = json.loads((ROOT / 'reports/consumer_availability_feasibility.json').read_text())
assert audit['spatial_mapping']['matched_ids'] == 2765
assert audit['weather_mapping']['matched_weather_hours'] == 695
sec = doc.sections[0]
sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.7)
sec.left_margin, sec.right_margin = Cm(1.9), Cm(1.9)
normal = doc.styles['Normal']
normal.font.size = Pt(11.5)
normal.paragraph_format.line_spacing = 1.25
normal.paragraph_format.space_after = Pt(9)
for p in sec.header.paragraphs:
    for r in p.runs: r.font.size = Pt(9)

def para(text, size=11.5, bold=False, after=9, color='252525'):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    base.run(p, text, size, bold, color)
    return p

def lead(text):
    p = para(text, 15, True, 15)
    p.paragraph_format.line_spacing = 1.22
    p.paragraph_format.keep_with_next = True
    return p

def h(text):
    p = para('□ ' + text, 12, True, 9)
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.keep_with_next = True

def point(label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    base.run(p, '○ ' + label + '  ', 11.5, True)
    base.run(p, text, 11.5)
    return p

def note(text): return para('※ ' + text, 9.5, after=8, color='626262')

def page(n, title):
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = n > 1
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(16)
    pp = p._p.get_or_add_pPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'DFDFDF'); pp.append(sh)
    borders = OxmlElement('w:pBdr')
    for side in ['top', 'bottom']:
        b = OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','5'),('color','777777'),('space','7')]: b.set(qn('w:'+k),v)
        borders.append(b)
    pp.append(borders)
    base.run(p, f'  {n:02d}  {title}', 13.5, True)

def table(headers, rows, widths, size=10.7):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c,w in zip(t.columns,widths): c.width=Cm(w)
    for i, values in enumerate([headers] + rows):
        cells = t.rows[0].cells if i == 0 else t.add_row().cells
        for j, text in enumerate(values):
            c=cells[j]; c.width=Cm(widths[j]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0]
            p.paragraph_format.space_before=Pt(8)
            p.paragraph_format.space_after=Pt(8)
            p.paragraph_format.line_spacing=1.2
            base.run(p,text,size,i==0)
            if i==0: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            cp=c._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E4E4E4' if i==0 else ('F6F6F6' if i%2 else 'FFFFFF')); cp.append(sh)
            bd=OxmlElement('w:tcBorders')
            for side in ['top','bottom','left','right']:
                b=OxmlElement('w:'+side)
                for k,v in [('val','single'),('sz','4'),('color','C8C8C8')]: b.set(qn('w:'+k),v)
                bd.append(b)
            cp.append(bd)
        cant=OxmlElement('w:cantSplit'); t.rows[i]._tr.get_or_add_trPr().append(cant)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    return t

def box(label,text): return table([label],[(text,)],[17.2],11.2)

page(1, '제안의 핵심')
p=para('따릉이, 어디로 빌리러 갈까?\n시민을 위한 대여소 선택 지원 서비스',18,True,15)
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=para('2026. 10. 19.  |  AI P-실무프로젝트 주제 검토 회의',10,after=15,color='626262')
p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
box('한 문장 제안', '현재 남은 자전거 수와 앞으로 0대가 될 위험을 함께 보여주어, 이용자가 빌리러 갈 대여소를 비교하고 선택하도록 돕는 서비스임.')
h('왜 필요한가')
point('현재 정보의 한계', '지금 1대가 있다는 사실만으로 이용 예정 시각에도 자전거가 남아 있을지는 알기 어려움.')
point('돕고 싶은 선택', '가까운 곳으로 갈지, 조금 더 떨어진 대안을 고를지 판단할 정보를 제공함.')
point('첫 사용자', '출근·등교 등 예정 시각이 있어 미리 대여소를 골라두려는 시민을 우선 대상으로 함.')
point('제공 방식', '휴대전화에서 여는 웹 화면으로 시작함. 지도에서 출발 위치를 고르고 이용 예정 시각을 입력하는 방식으로 시연함.')
h('이번 회의에서 판단할 사항')
para('“대여량을 맞히는 연구”에서 한 걸음 더 나아가,\n“빌리러 갈 곳을 고르는 서비스”로 추진할 가치가 있는가?')
note('필요성과 기대 효과는 기획 가설임. 실제 불편의 빈도와 기존 앱의 제공 기능은 사용자 인터뷰·기능 조사로 확인해야 함.')

page(2, '이용자는 이렇게 사용함')
lead('1시간 뒤 이용을 계획하는 상황부터 시연함')
para('예: 07시에 준비하면서 08시에 이용할 대여소를 고르는 상황.\n출발 위치와 이용 예정 시각을 입력하면 주변 대여소를 비교함.')
note('아래 수량·거리·위험 신호는 기능 설명을 위한 가상 예이며 실제 모델 결과가 아님.')
table(['비교 정보', 'A 대여소', 'B 대여소'],[
    ('출발지와의 직선거리','200m','450m'),
    ('조회 시점 수량','1대','6대'),
    ('목표 시각의 0대 위험','주의','상대적으로 낮음'),
    ('선택할 때 볼 점','가깝지만 다른 대안도 확인','조금 더 멀어도 함께 비교'),
],[4.6,6.3,6.3],11.2)
h('화면은 세 가지 정보를 중심으로 구성')
point('비교', '지도와 목록에서 현재 수량·거리·위험 신호를 함께 봄.')
point('근거', '자료 시각과 “수량이 줄어드는 흐름” 등의 참고 이유를 확인함. 설명 항목은 실제 모델 입력과 일치시킴.')
point('선택', '기본 추천 1곳과 주변 대안 최대 2곳을 비교하고 이용자가 결정함.')
note('직선거리는 실제 보행거리·시간과 다름. 실제 이용 시에는 최신 수량을 다시 확인하며, 서비스가 예약이나 대여를 보장하는 것은 아님.')

page(3, '어떤 데이터가 무엇을 뒷받침하는가')
lead('수량은 “예측”, 위치는 “선택”, 날씨는 “보조 설명”에 사용')
table(['데이터', '확인한 사실', '뒷받침하는 기능'],[
    ('D1  시간별 수량\n서울시 [1]', '2025년 1월 전체\n약 192만 행·대여소 2,766곳\n1시간 간격 기록 쌍 약 191만 개', '현재·최근 수량을 입력으로 사용.\n다음 기록이 0대인지 정답을 만들 수 있는 자료'),
    ('D2  대여소 위치\n서울시 [2]', '2024년 12월 위치와 연결\n2,766곳 중 2,765곳 연결\n2,711곳에 500m 이웃 존재', '출발지 주변 후보와 거리를 계산.\n가까운 대안이 있는지 비교'),
    ('D3  기상 관측\n기상청 [3]', '2025년 서울 연간 자료 확보\n수량 기록이 있는 695시간 모두\n같은 시각의 날씨와 연결', '날씨에 따라 흐름이 달라지는지 검토.\n예측 시점에 이용 가능한 값만 입력'),
],[3.5,6.8,6.9],10.7)
h('확인 결과의 의미')
point('할 수 있는 것', '공식 원본을 열고 번호·시간을 맞춰, 예측과 지도 비교에 필요한 자료가 연결되는지 확인했음.')
point('아직 확인할 것', '자료 연결이 된다는 사실은 정확도 향상을 뜻하지 않음. 운영 상태·수량의 의미·관측 시점 확인 후 본격 비교가 필요함.')
note('위 수치는 당시 자료 기준임. 위치 미연결 1곳은 번호 4721이며 별도 확인함. D3는 대표 관측소 1곳의 사후 관측본이므로 미래 날씨로 그대로 사용할 수 없음.')

page(4, '데이터가 추천으로 이어지는 과정')
lead('가깝다는 이유만으로 고르지 않고, 위험 신호를 함께 비교함')
table(['단계', '처리하는 내용', '사용 자료'],[
    ('① 주변 찾기', '출발지에서 직선거리 500m 안의 대여소를 후보로 찾음.', '위치 D2'),
    ('② 상태 확인', '시각이 확인된 수량을 붙임. 1대 이상인 후보를 우선 검토하며, 누락·오래된 정보는 구분함.', '수량 D1\n실서비스는 D4'),
    ('③ 위험 예측', '최근 수량 변화·요일·시간 등을 보고 한 시간 뒤 기록의 0대 위험을 계산함.', '수량 D1\n기상 D3는 선택'),
    ('④ 대안 제시', '위험 기준을 통과한 후보 중 가까운 곳을 기본 추천하고 대안도 제시함.', '위험 신호 + 거리'),
],[3.1,10.5,3.6],10.9)
h('판단 기준은 시험 전에 정함')
point('추천 조건', '500m 범위와 위험 판정 기준을 검증자료로 점검한 뒤 고정함. 적합한 후보가 없으면 무리한 추천 대신 확인 필요로 표시함.')
point('추천 이유', '“가까움”, “현재 수량 있음”, “목표 시각 0대 위험이 낮게 예상됨”을 구분해 보여줌.')
point('주변을 보는 이유', '같은 역·생활권에 있는 대여소의 변화가 예측에 도움이 될 수 있다는 가설임. 실제로 도움이 되는지는 자기 대여소만 보는 AI와 비교함.')
box('왜 AI까지 시험하는가', '현재 수량이 앞으로도 같다고 보는 방법보다, 수량의 변화와 주변 흐름을 함께 보는 방법이 더 잘 맞는지 검증하기 위함임. 개선이 없다면 AI의 추가 가치를 주장하지 않음.')

page(5, '지금 만들 수 있는 것과 추가 조건')
lead('과거 날짜로 검증하는 시연과 실시간 서비스를 구분함')
table(['1차 과제: 공개 과거 자료', '확장: 실제 이용 시각 안내'],[
    ('한 시간 뒤 기록의 0대 여부 예측\n과거 날짜의 대여소 비교·추천\n예상 결과와 실제 기록 확인', '공식 실시간 조회 연동\n더 짧은 간격의 수량 기록 확보\n5~10분 뒤 등 목표 시간별 재검증'),
],[8.6,8.6],11.2)
h('실서비스를 뒷받침할 추가 데이터')
point('D4  실시간 수량 [4]', '현재 수량·대여소 위치를 제공한다는 공식 안내는 확인함. 인증을 거친 실제 호출·갱신 지연·과거 이력 확보는 아직 검증하지 않았음.')
h('학습 전에 해결할 세 가지')
point('수량의 의미', '제공 페이지는 대여가능 수량이나 파일 열 이름은 “거치대수량”임. 실제 관측 방식·시각을 확인해야 함.')
point('누락·미운영', '1월에는 도시 전체 기록이 빠진 49시간이 있음. 이를 0대로 채우지 않으며 운영하지 않는 대여소도 구분함.')
point('시간 간격', '현재 1시간 자료는 5~10분 뒤 상황의 성능 근거가 될 수 없음. 짧은 예측은 촘촘한 이력 확보 후 확장함.')
point('기간 확대', '현재 전체 검사는 1월을 대상으로 했음. 겨울 한 달로 연중 효과를 주장하지 않도록 다른 달의 수량 자료를 추가 확보해 시험함.')
note('수량 0 기록은 약 10.1%이지만 이는 파일의 관측 비율임. 시민 대여 실패율·운영 중 품절 시간 비율 또는 헛걸음 감소 효과로 해석하지 않음.')

page(6, '무엇을 잘해야 좋은 서비스인가')
lead('“0대 위험을 놓치지 않는가”와 “추천이 도움이 되는가”를 평가')
para('정답: 현재에서 정확히 1시간 뒤 기록이 0대면 “위험”,\n1대 이상이면 “수량 있음”으로 구분함. 운영·관측 확인 범위에서 적용함.')
table(['지표', '이해하기 쉬운 뜻', '역할'],[
    ('F1\n0대 위험 기준', '위험이라고 알려준 것이 맞는지와 실제 위험을 놓치지 않았는지를 함께 봄.', '주평가지표\n클수록 좋음'),
    ('재현율·정밀도', '재현율: 실제 0대를 얼마나 미리 찾았는가\n정밀도: 위험 알림 중 실제 0대가 얼마나 되는가', '놓친 위험과\n불필요한 알림 확인'),
    ('추천 적중률\n과거 자료 기준', '추천한 대여소의 목표 기록에 1대 이상 있던 비율. 추천 보류율과 거리도 함께 봄.', '선택 지원의\n실질적 가치 확인'),
],[3.8,9.3,4.1],10.9)
h('정량 목표는 회의에서 검토할 제안임')
point('모델 목표', '기본 AI보다 F1 5%p 향상, 재현율 악화 없음을 제안함. 예를 들어 75%→80%는 5%p 향상이며, 실제 성능은 아직 없음.')
point('서비스 목표', '현재 1대 이상인 가장 가까운 곳을 고르는 방식보다 추천 적중률이 나아지는지 확인함.')
note('학교가 지정한 합격선이 아님. 최종 시험 전에 목표를 확정함. F1 = 2×정밀도×재현율÷(정밀도+재현율), 분모 0은 0 처리. 실제 0대 건수도 공개함.')
note('분류 학습은 잘못된 확신에 벌점을 주는 로그 손실을 사용함. L2·MAE·RMSE는 수량 자체를 예측할 때 별도로 검토함.')

page(7, '어떻게 공정하게 검증하고 평가에 연결할까')
table(['비교 방법', '비교하는 이유'],[
    ('① 지금과 같은 상태로 예상', '복잡한 AI 없이도 맞힐 수 있는 수준 확인'),
    ('② 기본 AI: 자기 수량·시간·가용 날씨', '자기 대여소 정보만 쓰는 충분한 기본 모델'),
    ('③ 같은 AI + 주변 수량 흐름', '주변 정보를 더한 효과만 분리하여 확인'),
],[8.2,9.0],10.7)
point('같은 시험', '과거로 학습하고 뒤의 기간으로 시험함. 같은 대여소·시간·정답·조정 조건을 쓰고, 당시 알 수 없는 날씨와 미래 기록은 제외함.')
point('추천도 같은 조건', '같은 가상 출발지·목표시각·후보로 비교함. 추천을 보류한 경우도 공개하고, 불리한 결과를 숨기지 않음.')
h('공식 평가 항목과 제출 근거')
table(['평가 항목', '이번 과제에서 보여줄 것'],[
    ('사전 40%: 적합성·창의성\n계획서·발표·질의응답', '시민의 선택 상황, 기존 기능 조사, 실제 데이터 연결, 비교 시험·역할·일정'),
    ('최종 60%: 성능·기본 방법 대비 우수성', '동일 조건 F1·재현율·정밀도와 추천 결과, 개선 및 실패 사례'),
    ('구현 완성도·적용 가능성', '입력→비교→선택 근거의 3개 절차 시연, 시민 사용성 의견'),
    ('문서·오픈소스·발표·동료평가', '출처·처리 규칙·재현 코드·GitHub/Wiki, 팀원별 결과와 기여'),
],[7.0,10.2],10.4)
note('근거: 첨부 평가 안내 2~3쪽 및 계획서 양식 3·6항. 실제 시민의 대여 성공·헛걸음 감소는 이 과거 자료 시험과 별도의 현장 검증이 필요함.')

page(8, '오늘 결정할 사항과 다음 행동')
lead('이용 장면·첫 개발 범위·검증 목표를 함께 확정')
point('주제', '시민이 빌리러 갈 대여소를 고르는 서비스로 진행할지 결정함.')
point('첫 범위', '1시간 뒤 예측을 과거 날짜로 시연하고, 실시간·짧은 도착시간 예측은 자료 확보 조건을 붙여 확장할지 결정함.')
table(['담당 영역', '사전평가 전 준비할 결과'],[
    ('본인: 예측·평가', '기본 방법과 개선 AI의 비교 설계, F1·추천 시험·목표치'),
    ('백엔드 경험자: 자료·서버', '시간·대여소별 자료 정리, 결과 전달 형식과 연결 계획'),
    ('팀원 C: 데이터 검증', '수량 의미·누락·운영 상태의 확인 결과와 근거 기록'),
    ('팀원 D: 화면·사용자', '대여소 비교 화면 시안, 불편·기존 기능 인터뷰 결과'),
],[4.2,13.0],10.7)
para('회의 결정: __________________________________________________\n담당·다음 점검일: _____________________________________________',10.5,after=9)
note('사용자 조사는 참여자 3~5명에게 불편 경험과 화면 선택 이유를 묻는 소규모 탐색부터 제안함. 인구 전체의 효과로 일반화하지 않음.')
para('사전 제출  10/25(일) 23:59  /  최종 제출  12/13(일) 23:59',10.5,True,9)
note('사전: PPT·PDF·10분 이내 영상 주소 TXT·계획서 WORD를 ZIP으로 팀장 1명 제출. [1]~[4]는 아래 제목을 누르면 공식 출처로 이동함.')
refs=[
    ('[1] 서울시 시간별 대여가능 수량','https://data.seoul.go.kr/dataList/OA-22382/F/1/datasetView.do'),
    ('[2] 서울시 대여소 위치','https://data.seoul.go.kr/dataList/OA-13252/F/1/datasetView.do'),
    ('[3] 기상청 ASOS','https://data.kma.go.kr/data/grnd/selectAsosList.do?pgmNo=36'),
    ('[4] 서울시 실시간 대여정보','https://data.seoul.go.kr/dataList/OA-15493/A/1/datasetView.do'),
]
for i in range(0,4,2):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(4)
    for j,(label,url) in enumerate(refs[i:i+2]):
        if j: base.run(p,'  |  ',9.5,color='626262')
        base.hyperlink(p,label,url)
        hl=p._p[-1]
        for sz in hl.findall('.//'+qn('w:sz')): sz.set(qn('w:val'),'19')
note('검사 기록: reports/consumer_availability_feasibility.json. 확인일 2026.10.19. 이번 자료는 기획·데이터 확인 결과이며 모델 성능·실시간 연결 완료 보고서가 아님.')

page(9, '4명이 나눠 맡을 일과 협업 방식')
lead('각자 작동하는 결과물을 맡고, 자료 형식을 맞춰 연결함')
table(['담당', '주요 업무', '책임질 결과물'],[
    ('A  본인\n모델·성능 평가', '과거 상태 기준·기본 AI·주변 정보 AI 비교. 기간 분리와 F1·재현율·추천 평가를 설계하고 학습·분석.', '예측 실행 코드\n성능 비교표\n대여소·시간별 예측 파일'),
    ('B  백엔드 경험자\n자료 처리·서버', '수량·위치·날씨 수집과 연결. 예측 파일을 읽어 화면에 전달. 실행 환경과 전체 연결 담당. 실시간 조회 가능 여부 확인.', '자료 처리 코드\n결과 조회 기능\n연결된 실행 환경'),
    ('C  팀원\n데이터 검증', '제공기관 설명·수량 의미·자료 시각·미운영 여부 확인. 검증 스크립트를 실행·수정해 누락·중복·위치 연결 점검.', '검증 스크립트·결과표\n출처·처리 규칙\n자료 오류·확인 이력'),
    ('D  팀원\n화면·사용자 검증', '간단한 웹 화면에서 위치·시간 입력, 지도·대여소 비교 구현. 시민 인터뷰와 화면 사용성 점검. 시연 과정 구성.', '작동하는 화면\n사용자 의견 기록\n시연·화면 설명 자료'),
],[3.7,8.3,5.2],10.4)
h('서로 주고받을 결과를 먼저 정함')
point('자료 → 예측 → 화면', 'B가 정리한 자료를 C가 점검하고 A가 학습함. A는 대여소 번호·예측 대상 시각·위험 신호·자료 시각을 결과표로 전달하고 B가 D의 화면에 연결함.')
point('기다리지 않고 시작', 'D는 합의한 형식의 가상 예시로 화면부터 만들고 나중에 실제 예측으로 교체함. C는 초보자도 확인할 수 있는 체크리스트부터 시작하며, 정답 규칙은 A·B와 함께 확정함.')
note('비전공 팀원 2명도 각각 검증 코드와 화면이라는 구현 결과를 갖도록 배분한 추천안임. 계획서·발표는 각자 담당 부분을 작성하고 함께 검토함. 공동 작업과 개인 기여를 기록하여 동료평가에 연결함.')

props=doc.core_properties
props.title='따릉이, 어디로 빌리러 갈까? — 시민용 서비스 구체화 및 4인 역할분담 회의자료'
props.subject='시민 사용 장면·데이터별 근거·추천 규칙·평가·4인 역할분담'
props.author='P_Project'
props.created=props.modified=datetime(2026,10,19,tzinfo=timezone.utc)
path=OUT/'따릉이_시민용서비스_회의자료_4인역할분담_20261019.docx'
doc.save(path)
print(path)
