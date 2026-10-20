"""Create the baseline, before/after, and official-rubric revision of the handout."""
from datetime import datetime, timezone
import json
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import meeting_doc_format as base

doc, OUT, ROOT = base.doc, base.OUT, base.ROOT
audit = json.loads((ROOT / 'reports/consumer_availability_feasibility.json').read_text())
assert audit['spatial_mapping']['matched_ids'] == 2765
assert audit['weather_mapping']['matched_weather_hours'] == 695
sec = doc.sections[0]
sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.7)
sec.left_margin, sec.right_margin = Cm(1.9), Cm(1.9)
normal = doc.styles['Normal']
normal.font.size = Pt(11.5)
normal.paragraph_format.line_spacing = 1.25
normal.paragraph_format.space_after = Pt(9)
for p in sec.header.paragraphs:
    for r in p.runs: r.font.size = Pt(9)

def para(text, size=11.5, bold=False, after=9, color='252525'):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    base.run(p, text, size, bold, color)
    return p

def lead(text):
    p = para(text, 15, True, 15)
    p.paragraph_format.line_spacing = 1.22
    p.paragraph_format.keep_with_next = True
    return p

def h(text):
    p = para('□ ' + text, 12, True, 9)
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.keep_with_next = True

def point(label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    base.run(p, '○ ' + label + '  ', 11.5, True)
    base.run(p, text, 11.5)
    return p

def note(text): return para('※ ' + text, 9.5, after=8, color='626262')

def page(n, title):
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = n > 1
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(16)
    pp = p._p.get_or_add_pPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'DFDFDF'); pp.append(sh)
    borders = OxmlElement('w:pBdr')
    for side in ['top', 'bottom']:
        b = OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','5'),('color','777777'),('space','7')]: b.set(qn('w:'+k),v)
        borders.append(b)
    pp.append(borders)
    base.run(p, f'  {n:02d}  {title}', 13.5, True)

def table(headers, rows, widths, size=10.7, pad=8):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c,w in zip(t.columns,widths): c.width=Cm(w)
    for i, values in enumerate([headers] + rows):
        cells = t.rows[0].cells if i == 0 else t.add_row().cells
        for j, text in enumerate(values):
            c=cells[j]; c.width=Cm(widths[j]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0]
            p.paragraph_format.space_before=Pt(pad)
            p.paragraph_format.space_after=Pt(pad)
            p.paragraph_format.line_spacing=1.2
            base.run(p,text,size,i==0)
            if i==0: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            cp=c._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E4E4E4' if i==0 else ('F6F6F6' if i%2 else 'FFFFFF')); cp.append(sh)
            bd=OxmlElement('w:tcBorders')
            for side in ['top','bottom','left','right']:
                b=OxmlElement('w:'+side)
                for k,v in [('val','single'),('sz','4'),('color','C8C8C8')]: b.set(qn('w:'+k),v)
                bd.append(b)
            cp.append(bd)
        cant=OxmlElement('w:cantSplit'); t.rows[i]._tr.get_or_add_trPr().append(cant)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    return t

def box(label,text): return table([label],[(text,)],[17.2],11.2)

