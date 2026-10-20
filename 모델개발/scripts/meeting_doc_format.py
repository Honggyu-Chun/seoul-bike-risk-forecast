"""Generate the plain-language handout with rationale, explicit metrics and rubric mapping."""
from pathlib import Path
from datetime import datetime, timezone
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables'
OUT.mkdir(exist_ok=True)
FONT = 'Noto Sans CJK KR'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.65), Cm(1.55)
sec.left_margin, sec.right_margin = Cm(1.8), Cm(1.8)
sec.header_distance, sec.footer_distance = Cm(.65), Cm(.7)

def run(p, text, size=10.5, bold=False, color='222222'):
    r = p.add_run(text)
    r.font.name = FONT
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    r.font.size = Pt(size)
    r.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    return r

for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    s = doc.styles[name]
    s.font.name = FONT
    s.font.size = Pt(10.5)
    s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    s.paragraph_format.line_spacing = 1.18
    s.paragraph_format.space_after = Pt(5)
    s.font.color.rgb = RGBColor.from_string('222222')
header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run(header, 'AI P-실무프로젝트  |  주제 선정 회의자료', 8, color='666666')
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(footer, '- ', 8, color='666666')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
run(footer, ' -', 8, color='666666')

def para(text, size=10.5, bold=False, color='222222', after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    run(p, text, size, bold, color)
    return p

def bullet(label, text, sub=False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(.65 if sub else .15)
    p.paragraph_format.first_line_indent = Cm(-.3 if sub else -.15)
    p.paragraph_format.space_after = Pt(5)
    run(p, '- ' if sub else '○ ')
    if label:
        run(p, label + ' ', bold=True)
    run(p, text)
    return p

def note(text): return para('※ ' + text, 8.5, color='666666', after=4)

def section(num, title, new_page=False):
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = new_page
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(11)
    pp = p._p.get_or_add_pPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'DEDEDE'); pp.append(sh)
    borders = OxmlElement('w:pBdr')
    for side in ['top','bottom']:
        el = OxmlElement('w:' + side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'5'); el.set(qn('w:color'),'777777'); el.set(qn('w:space'),'6'); borders.append(el)
    pp.append(borders)
    run(p, f'  {num}.  {title}', 13, True)

def sub(title):
    p = para('□ ' + title, 11, True, after=6)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.keep_with_next = True

def table(head, rows, widths, size=9.3):
    t = doc.add_table(rows=1, cols=len(head))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col,w in zip(t.columns,widths): col.width = Cm(w)
    for row_i, values in enumerate([head]+rows):
        cells = t.rows[0].cells if row_i == 0 else t.add_row().cells
        for j, text in enumerate(values):
            c=cells[j]; c.width=Cm(widths[j]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0]
            p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(4)
            p.paragraph_format.line_spacing=1.12
            run(p,text,size,row_i==0)
            if row_i==0: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            cp=c._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E5E5E5' if row_i==0 else ('F5F5F5' if row_i%2 else 'FFFFFF')); cp.append(sh)
            bd=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'C5C5C5'); bd.append(e)
            cp.append(bd)
        cant=OxmlElement('w:cantSplit'); t.rows[row_i]._tr.get_or_add_trPr().append(cant)
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    return t

def box(title,text):
    t=table([title],[(text,)],[17.4],10)
    t.rows[0].cells[0].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.LEFT
    return t

def hyperlink(p,label,url):
    rel=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rel)
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    for tag,attr,value in [('color','val','444444'),('sz','val','16'),('u','val','single')]:
        el=OxmlElement('w:'+tag); el.set(qn('w:'+attr),value); pr.append(el)
    rf=OxmlElement('w:rFonts'); rf.set(qn('w:ascii'),FONT); rf.set(qn('w:eastAsia'),FONT); pr.append(rf)
    r.append(pr); tx=OxmlElement('w:t'); tx.text=label; r.append(tx); h.append(r); p._p.append(h)

