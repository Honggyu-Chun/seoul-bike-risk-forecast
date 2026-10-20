"""Generate a grayscale meeting paper for readers without an AI background."""
from pathlib import Path
from datetime import datetime, timezone
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables'
OUT.mkdir(exist_ok=True)
FONT = 'Noto Sans CJK KR'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.65), Cm(1.55)
sec.left_margin, sec.right_margin = Cm(1.8), Cm(1.8)
sec.header_distance, sec.footer_distance = Cm(.65), Cm(.7)

def run(p, text, size=10.5, bold=False, color='222222'):
    r = p.add_run(text)
    r.font.name = FONT
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    r.font.size = Pt(size)
    r.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    return r

for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    s = doc.styles[name]
    s.font.name = FONT
    s.font.size = Pt(10.5)
    s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    s.paragraph_format.line_spacing = 1.18
    s.paragraph_format.space_after = Pt(5)
    s.font.color.rgb = RGBColor.from_string('222222')
header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run(header, 'AI P-실무프로젝트  |  주제 선정 회의자료', 8, color='666666')
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(footer, '- ', 8, color='666666')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
run(footer, ' -', 8, color='666666')

def para(text, size=10.5, bold=False, color='222222', after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    run(p, text, size, bold, color)
    return p

def bullet(label, text, sub=False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(.65 if sub else .15)
    p.paragraph_format.first_line_indent = Cm(-.3 if sub else -.15)
    p.paragraph_format.space_after = Pt(5)
    run(p, '- ' if sub else '○ ')
    if label:
        run(p, label + ' ', bold=True)
    run(p, text)
    return p

def note(text): return para('※ ' + text, 8.5, color='666666', after=4)

def section(num, title, new_page=False):
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = new_page
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(11)
    pp = p._p.get_or_add_pPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'DEDEDE'); pp.append(sh)
    borders = OxmlElement('w:pBdr')
    for side in ['top','bottom']:
        el = OxmlElement('w:' + side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'5'); el.set(qn('w:color'),'777777'); el.set(qn('w:space'),'6'); borders.append(el)
    pp.append(borders)
    run(p, f'  {num}.  {title}', 13, True)

def sub(title):
    p = para('□ ' + title, 11, True, after=6)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.keep_with_next = True

def table(head, rows, widths, size=9.3):
    t = doc.add_table(rows=1, cols=len(head))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col,w in zip(t.columns,widths): col.width = Cm(w)
    for row_i, values in enumerate([head]+rows):
        cells = t.rows[0].cells if row_i == 0 else t.add_row().cells
        for j, text in enumerate(values):
            c=cells[j]; c.width=Cm(widths[j]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0]
            p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(4)
            p.paragraph_format.line_spacing=1.12
            run(p,text,size,row_i==0)
            if row_i==0: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            cp=c._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E5E5E5' if row_i==0 else ('F5F5F5' if row_i%2 else 'FFFFFF')); cp.append(sh)
            bd=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'C5C5C5'); bd.append(e)
            cp.append(bd)
        cant=OxmlElement('w:cantSplit'); t.rows[row_i]._tr.get_or_add_trPr().append(cant)
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    return t

def box(title,text):
    t=table([title],[(text,)],[17.4],10)
    t.rows[0].cells[0].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.LEFT
    return t

def hyperlink(p,label,url):
    rel=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rel)
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    for tag,attr,value in [('color','val','444444'),('sz','val','16'),('u','val','single')]:
        el=OxmlElement('w:'+tag); el.set(qn('w:'+attr),value); pr.append(el)
    rf=OxmlElement('w:rFonts'); rf.set(qn('w:ascii'),FONT); rf.set(qn('w:eastAsia'),FONT); pr.append(rf)
    r.append(pr); tx=OxmlElement('w:t'); tx.text=label; r.append(tx); h.append(r); p._p.append(h)

p=para('따릉이 이용량 예측 서비스 검토(안)',18,True,after=9)
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=para('2026. 10. 19.  /  AI P-실무프로젝트 주제 검토',9.5,color='666666',after=10)
p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
box('무엇을 하려는가', '“언제, 어느 대여소에서 따릉이가 몇 건 대여될지”를 예상하여 서울 지도에 보여주는 서비스를 만들고자 함. 담당자는 이용이 많을 것으로 예상되는 곳을 먼저 살펴볼 수 있음. 필요한 주요 자료는 실제 파일로 확보하였으며, 얼마나 잘 맞히는지는 앞으로 시험할 예정임.')

section('1','왜 필요한가')
sub('가. 지금 남은 자전거 수와 앞으로 빌려갈 건수는 다른 정보임')
bullet('현재 확인 가능한 정보', '서울시 공개 서비스는 대여소 위치와 지금 빌릴 수 있는 자전거 수 등을 제공함.[1]')
bullet('추가로 알고 싶은 정보', '운영 담당자에게는 “현재 몇 대가 있는가”와 함께 “앞으로 어느 곳의 이용이 많을 것인가”도 필요함. 두 정보를 구분하여 볼 수 있도록 하는 것이 이번 제안의 출발점임.')
box('이해를 돕기 위한 가상 사례', '현재 자전거가 10대씩 있는 두 대여소가 있다고 가정함. 한 곳은 다음 한 시간에 2건, 다른 곳은 15건의 대여가 예상된다면 담당자가 관심 있게 볼 곳은 달라질 수 있음. 단, 실제 부족 여부는 반납과 자전거 이동까지 함께 확인해야 함.')
sub('나. 서울 전역을 살펴보는 운영 담당자를 위한 도구임')
bullet('사용자', '대여소 운영을 살펴보는 담당자를 우선 대상으로 함. 시민이 자전거를 빌릴 때 사용하는 앱과는 사용 목적이 다름.')
bullet('사용 목적', '많은 대여소 중에서 이용이 많을 시간과 장소를 한눈에 확인하고, 점검할 곳을 고르는 데 참고하도록 함.')
bullet('기대 효과', '대여소별 자료를 일일이 찾아보는 부담을 줄이고, 같은 기준으로 비교할 수 있음. 실제 업무가 얼마나 편해지는지는 화면 시연과 담당자 의견으로 확인할 필요가 있음.')
sub('다. 이번 과제에서는 “대여 건수 예측”부터 시작함')
bullet('목표', '대여소별로 한 시간 동안 몇 건이 대여될지 계산하고, 실제 대여 건수와 비교하여 잘 맞는지 확인함. 공연·행사 정보는 반드시 넣지 않아도 됨.')
note('자전거가 없어 빌리지 못한 사람 수나 자전거를 어디로 옮겨야 하는지까지 계산하는 과제는 아님. 이 문서의 “이용량”은 실제 대여가 이루어진 건수를 뜻함.')

section('2','기존 방식과 무엇이 다른가',True)
sub('가. 현재 수량 조회에 더해, 시간대별 예상 이용량을 보여줌')
table(['구분','주로 답하는 질문','확인하거나 만들 내용'],[
    ('서울시 공개 대여정보 [1]','“지금 어디에서 몇 대를 빌릴 수 있는가?”','현재 대여가능 수량과 위치를 확인하는 정보'),
    ('이번에 제안하는 서비스','“어느 시간에 어느 대여소의 이용이 많을 것인가?”','예상 대여 건수를 지도에 표시하고, 과거 예측이 얼마나 맞았는지 함께 제시'),
],[3.7,6.7,7.0],9.7)
note('공식 공개 안내에서 확인한 기능과 비교한 것임. 운영기관 내부에서 사용하는 예측 방식까지 조사한 것은 아니므로, 기존에 예측 기능이 전혀 없다고 단정하지 않음.')
sub('나. 해당 대여소의 기록뿐 아니라 주변 이용 흐름도 참고함')
bullet('기본적인 계산', 'A대여소가 예전에 같은 요일·시간에 얼마나 이용되었는지, 그때의 날씨와 휴일 여부는 어땠는지 참고하여 예상 건수를 계산함.')
bullet('추가로 시험할 계산', 'A대여소 주변이나 이용 흐름이 비슷한 다른 대여소의 과거 기록도 함께 참고함. “이곳만 볼 때보다 주변도 함께 보면 더 잘 맞는가”를 확인하는 것임.')
bullet('이렇게 생각한 이유', '같은 지하철역이나 업무지역 주변의 대여소는 비슷한 시간에 이용이 늘어날 수 있음. 실제로 도움이 되는지는 같은 조건의 비교 시험으로 확인해야 함.')
bullet('주장할 차이', '새로운 AI 기술을 발명했다고 하기보다, 주변 기록을 더했을 때 어느 곳에서 예측이 나아지는지 확인하고 그 결과를 보기 쉬운 화면으로 제공하는 데 의미를 둠.')
sub('다. 담당자는 다음 순서로 사용하게 됨')
table(['순서','화면에서 하는 일'],[
    ('① 날짜·시간 선택','살펴볼 날짜와 시간대를 고르면 서울 지도에 대여소별 예상 대여 건수가 표시됨.'),
    ('② 관심 대여소 선택','대여소를 누르면 하루 중 어느 시간에 이용이 많을지 그래프로 확인함.'),
    ('③ 지난 예측 확인','과거 날짜를 선택하면 예상 건수와 실제 건수를 나란히 보고, 얼마나 틀렸는지 확인함.'),
],[4.0,13.4],9.7)
note('위 화면은 앞으로 만들 기능의 구상임. 서울시 대여 건수 기록은 반년 단위로 갱신되므로 첫 시연은 과거 날짜의 예측과 실제 결과를 비교하는 방식으로 시작함. 현재 상황을 즉시 반영하려면 최신 이용 기록을 별도로 받아야 함.')

section('3','실제로 쓸 자료가 있는가',True)
sub('가. 핵심 자료 3종을 실제로 내려받아 확인함')
table(['자료·출처','무엇을 확인했는가','어디에 쓰는가'],[
    ('따릉이 이용 기록\n서울시 [2]','2025년 1월 기록을 확인함.\n대여소 번호 2,730개가 등장하며,\n실제 대여 건수 합계는 약 170만 건임.','어느 대여소에서 몇 시에 몇 건 대여되었는지 확인하고 예측을 연습·평가하는 데 사용'),
    ('대여소 위치\n서울시 [3]','2024년 12월 위치 자료와 대조함.\n2,730개 중 2,728개의 위치를 찾음.\n이 중 2,674곳은 500m 안에 이웃 대여소가 있음.','지도에 대여소를 표시하고, 가까운 대여소끼리 연결하는 데 사용'),
    ('기온·습도·바람 등\n기상청 [4]','서울 대표 관측소 1곳의 2025년 자료를 확보함. 1월의 모든 시간대에 해당하는 날씨가 있고, 기온·습도·풍속 값은 빠짐없이 있음.','날씨에 따라 이용량이 어떻게 달라지는지 확인. 관측소 1곳으로 서울을 대표하는 한계가 있음'),
],[3.7,7.8,5.9],9.5)
sub('나. 사이트에 있다는 사실만 확인한 것이 아니라 내용을 대조함')
bullet('파일 열기', '공식 기관의 파일을 내려받아 읽을 수 있는지, 필요한 날짜·시간·대여소 번호·이용 건수가 있는지 확인함.')
bullet('건수 계산', '이용 기록은 나이·이용권 종류 등에 따라 여러 줄로 나뉘어 있음. 줄 수를 사람 수로 세지 않고, 파일에 적힌 “이용건수”를 더함. 합계는 1,704,190건임.')
bullet('위치·날씨 연결', '대여소 번호로 위치를 찾고, 같은 날짜·시간의 이용 기록과 날씨를 묶음. 1월 31일×24시간인 744개 시간대에 날씨가 모두 연결됨.')
note('확인 결과는 자료를 함께 쓸 수 있다는 뜻임. 모든 대여소가 매시간 빠짐없이 기록되었다거나, AI가 이미 정확해졌다는 뜻은 아님. 대여소 수는 당시 자료 기준임.')
sub('다. 추가로 찾아본 자료와 남은 준비')
bullet('지하철·공원', '공식 자료를 실제로 열어 확인함.[6][7][8] 역별 승하차 인원이나 공원과의 거리를 추가할 수 있으나, 첫 단계의 필수 자료로 정하지는 않음.')
bullet('공휴일·계절', '최신 2025년 공식 공휴일 목록은 추가 확보가 필요함.[9] 또 겨울 한 달로 1년 전체를 평가할 수 없으므로 다른 계절의 이용 기록을 더 확보해야 함.')
note('해외 대학의 서울 따릉이 공개 자료도 확인했으나 대여소 번호가 없어, 대여소별 지도를 만드는 주자료로는 사용하지 않음.[5]')

section('4','예측이 나아졌는지 어떻게 확인하는가',True)
sub('가. 같은 대여소·같은 기간을 세 가지 방법으로 계산함')
table(['비교 방법','쉽게 설명하면','확인할 내용'],[
    ('① 과거 평균 사용','“예전 같은 요일·시간에 평균 8건 이용되었으니, 이번에도 8건일 것”으로 계산','복잡한 계산 없이도 맞힐 수 있는 수준'),
    ('② 기본 AI 사용','과거 이용 기록, 요일·휴일, 사용할 수 있는 날씨를 보고 AI가 예상 건수를 계산','과거 평균보다 잘 맞는지 확인'),
    ('③ 주변 기록까지 참고','②와 같은 AI에 주변이나 이용 흐름이 비슷한 대여소의 과거 기록을 추가','주변 기록을 더한 것이 도움이 되는지 확인'),
],[3.6,8.3,5.5],9.7)
note('여기서 AI는 과거 자료에서 이용 규칙을 찾아 예상 건수를 계산하는 프로그램을 뜻함. “모델”이라는 말도 이 계산 방식을 가리킴.')
sub('나. “평균 몇 건 틀렸는가”로 결과를 설명함')
bullet('평가 방법', '예상 건수와 실제 건수 사이의 차이가 몇 건인지 대여소별·시간대별로 구해 평균을 냄. 평균 차이가 작을수록 더 잘 맞는 것으로 판단함.')
bullet('설명용 예', '실제 10건인데 8건으로 예상했다면 2건 틀린 것임. 이런 차이를 여러 시간대에 걸쳐 계산함. 이 숫자는 설명을 위한 예이며 실제 시험 결과가 아님.')
bullet('함께 볼 결과', '서울 전체 평균뿐 아니라 출퇴근 시간과 이용이 많은 대여소도 따로 확인함. 크게 틀린 사례와 개선되지 않은 곳도 함께 제시함.')
sub('다. 답을 미리 보고 맞히는 시험이 되지 않도록 함')
bullet('시간 순서 준수', '먼저 공개된 과거 자료로 계산 방법을 익힌 뒤, 그보다 뒤의 기간을 맞혀 보게 함. 세 방법 모두 같은 대여소와 기간으로 비교함.')
bullet('미래 정보 제외', '예측하는 시점에 알 수 없었던 실제 날씨나 이후의 대여 기록은 사용하지 않음. AI를 조정하는 조건도 같게 하여 주변 기록의 효과만 비교함.')
bullet('빠진 기록 확인', '기록이 없는 칸을 무조건 “이용 0건”으로 보지 않음. 실제로 아무도 빌리지 않은 것인지, 자료가 빠졌거나 운영하지 않은 것인지 확인한 뒤 처리함.')
box('현재까지 확인한 것과 앞으로 할 것', '자료 확보와 위치·시간 연결은 완료함. 세 가지 방법의 본격적인 정확도 비교와 화면 제작은 앞으로 할 일임. 따라서 지금은 “몇 % 더 정확하다”거나 “자전거 부족이 줄었다”고 결론 내릴 단계가 아님.')

section('5','과제 평가와 오늘 회의에서 정할 사항',True)
sub('가. 평가자에게 보여줄 결과')
table(['평가 구분','평가자가 확인할 내용','준비할 자료'],[
    ('사전평가 40%','주제가 필요한가, 기존 방식과 차이가 있는가, 실현 가능한 계획인가','필요성·사용 장면, 실제 자료 확인 결과, 비교 시험 계획, 역할·일정과 발표자료'),
    ('최종평가 60%','기본 방법보다 잘 맞는가, 화면이 작동하는가, 실제로 활용할 수 있는가','정확도 비교표, 작동하는 지도·그래프, 자료 출처·재현 방법, 공개 코드와 팀원별 기여 기록'),
],[3.0,6.5,7.9],9.3)
note('첨부 평가 안내 2~3쪽 기준임. 최종평가에는 발표·질의응답과 동료평가도 포함됨.')
sub('나. 오늘 회의에서 합의할 세 가지')
bullet('주제', '“서울 대여소별 예상 대여 건수를 보여주는 서비스”로 진행할지, 주변 대여소 기록의 효과를 핵심 비교 내용으로 삼을지 결정함.')
bullet('첫 개발 범위', '지도·예상 건수 그래프·과거 실제값 비교부터 만들지 결정함. 지하철·공원 정보와 현재 상황의 즉시 반영은 후속 단계로 둘 수 있음.')
bullet('담당', '① 자료 정리·확인 ② 예측 계산·정확도 비교 ③ 계산 결과를 화면에 전달 ④ 화면·발표·문서의 담당자와 점검일을 정함.')
para('회의 결정: __________________________________________________________\n담당 및 다음 점검일: __________________________________________________',9.5,after=5)
sub('다. 가까운 제출 일정')
para('10/19~21  주제·개발 범위·담당 확정, 추가로 필요한 자료 정리\n10/22~24  계획서·발표자료·화면 예시 작성 및 검토\n10/25(일) 23:59  사전평가 제출: PPT·PDF·10분 이내 영상의 주소 TXT·계획서 WORD를 압축하여 팀장 1명이 제출\n12/13(일) 23:59  최종평가 제출: 정확도 비교·서비스·결과보고서 등 정리',9,after=5)
note('공휴일 보완은 [9]의 “공휴일 정보 조회”에서 2025년 전체 목록을 받으면 됨. 나중에 지정된 1/27·6/3도 포함되었는지 확인하며, 인증키가 아닌 조회 결과 파일만 공유함.')
sub('참고 자료  ※ 제목 클릭 시 공식 출처로 이동')
refs=[
 ('[1] 서울시 실시간 대여정보','https://data.seoul.go.kr/dataList/OA-15493/A/1/datasetView.do'),
 ('[2] 시간대별 이용정보','https://data.seoul.go.kr/dataList/OA-15245/F/1/datasetView.do'),
 ('[3] 대여소 위치','https://data.seoul.go.kr/dataList/OA-13252/F/1/datasetView.do'),
 ('[4] 기상청 ASOS','https://data.kma.go.kr/data/grnd/selectAsosList.do?pgmNo=36'),
 ('[5] UCI Seoul Bike','https://archive.ics.uci.edu/dataset/560/seoul+bike+sharing+demand'),
 ('[6] 지하철 일별 승하차','https://data.seoul.go.kr/dataList/OA-12914/F/1/datasetView.do'),
 ('[7] 지하철 시간대별 월 누계','https://data.seoul.go.kr/dataList/OA-12252/S/1/datasetView.do'),
 ('[8] 서울 주요 공원','https://data.seoul.go.kr/dataList/OA-394/S/1/datasetView.do'),
 ('[9] 한국천문연구원 특일 정보','https://www.data.go.kr/data/15012690/openapi.do'),
]
for i in range(0,len(refs),3):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3)
    for j,(label,url) in enumerate(refs[i:i+3]):
        if j: run(p,'  |  ',8,color='666666')
        hyperlink(p,label,url)

props=doc.core_properties
props.title='따릉이 이용량 예측 서비스 검토(안) — 쉬운 설명본'
props.subject='서비스 필요성·사용 장면·데이터 확인·정확도 비교·회의 결정 사항'
props.author='P_Project'
props.keywords='따릉이, 추진 검토, 회의자료, 데이터 검증, 수요 예측'
props.created=props.modified=datetime(2026,10,19,tzinfo=timezone.utc)
path=OUT/'따릉이_이용량예측_회의자료_쉬운설명본_20261019.docx'
doc.save(path)
print(path)
