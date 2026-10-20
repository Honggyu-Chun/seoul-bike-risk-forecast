"""Generate the revised, grayscale Korean meeting paper with a problem-first argument."""
from pathlib import Path
from datetime import datetime, timezone
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables'
OUT.mkdir(exist_ok=True)
FONT = 'Noto Sans CJK KR'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.65), Cm(1.55)
sec.left_margin, sec.right_margin = Cm(1.8), Cm(1.8)
sec.header_distance, sec.footer_distance = Cm(.65), Cm(.7)

def run(p, text, size=10.5, bold=False, color='222222'):
    r = p.add_run(text)
    r.font.name = FONT
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    r.font.size = Pt(size)
    r.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    return r

for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    s = doc.styles[name]
    s.font.name = FONT
    s.font.size = Pt(10.5)
    s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    s.paragraph_format.line_spacing = 1.18
    s.paragraph_format.space_after = Pt(5)
    s.font.color.rgb = RGBColor.from_string('222222')
header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run(header, 'AI P-실무프로젝트  |  주제 선정 회의자료', 8, color='666666')
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(footer, '- ', 8, color='666666')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
run(footer, ' -', 8, color='666666')

def para(text, size=10.5, bold=False, color='222222', after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    run(p, text, size, bold, color)
    return p

def bullet(label, text, sub=False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(.65 if sub else .15)
    p.paragraph_format.first_line_indent = Cm(-.3 if sub else -.15)
    p.paragraph_format.space_after = Pt(5)
    run(p, '- ' if sub else '○ ')
    if label:
        run(p, label + ' ', bold=True)
    run(p, text)
    return p

def note(text): return para('※ ' + text, 8.5, color='666666', after=4)

def section(num, title, new_page=False):
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = new_page
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(11)
    pp = p._p.get_or_add_pPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'DEDEDE'); pp.append(sh)
    borders = OxmlElement('w:pBdr')
    for side in ['top','bottom']:
        el = OxmlElement('w:' + side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'5'); el.set(qn('w:color'),'777777'); el.set(qn('w:space'),'6'); borders.append(el)
    pp.append(borders)
    run(p, f'  {num}.  {title}', 13, True)

def sub(title):
    p = para('□ ' + title, 11, True, after=6)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.keep_with_next = True

def table(head, rows, widths, size=9.3):
    t = doc.add_table(rows=1, cols=len(head))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col,w in zip(t.columns,widths): col.width = Cm(w)
    for row_i, values in enumerate([head]+rows):
        cells = t.rows[0].cells if row_i == 0 else t.add_row().cells
        for j, text in enumerate(values):
            c=cells[j]; c.width=Cm(widths[j]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0]
            p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(4)
            p.paragraph_format.line_spacing=1.12
            run(p,text,size,row_i==0)
            if row_i==0: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            cp=c._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E5E5E5' if row_i==0 else ('F5F5F5' if row_i%2 else 'FFFFFF')); cp.append(sh)
            bd=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'C5C5C5'); bd.append(e)
            cp.append(bd)
        cant=OxmlElement('w:cantSplit'); t.rows[row_i]._tr.get_or_add_trPr().append(cant)
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    return t

def box(title,text):
    t=table([title],[(text,)],[17.4],10)
    t.rows[0].cells[0].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.LEFT
    return t

def hyperlink(p,label,url):
    rel=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rel)
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    for tag,attr,value in [('color','val','444444'),('sz','val','16'),('u','val','single')]:
        el=OxmlElement('w:'+tag); el.set(qn('w:'+attr),value); pr.append(el)
    rf=OxmlElement('w:rFonts'); rf.set(qn('w:ascii'),FONT); rf.set(qn('w:eastAsia'),FONT); pr.append(rf)
    r.append(pr); tx=OxmlElement('w:t'); tx.text=label; r.append(tx); h.append(r); p._p.append(h)

p=para('서울시 공공자전거 대여량 예측 및\n운영지원 서비스 추진 검토(안)',18,True,after=9)
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=para('2026. 10. 19.  /  AI P-실무프로젝트 주제 검토',9.5,color='666666',after=12)
p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
box('검토 요지', '서울 전역 대여소의 시간대별 대여량을 예측하여 운영 검토에 활용하는 방안임. 핵심 데이터의 확보·연결 가능성은 확인하였으며, 주변 대여소 정보를 활용한 정확도 개선 여부는 후속 실험으로 검증하고자 함.')

section('1','추진 배경 및 필요성')
sub('가. 현재 정보만으로는 앞으로의 이용 규모를 파악하기 어려움')
bullet('현황 정보의 역할', '서울시 실시간 대여정보는 대여소 위치, 대여가능 수량, 거치율 등 현재 상태를 제공함.[1]')
bullet('추가로 필요한 정보', '운영 담당자가 특정 시간대의 이용 집중을 미리 검토하려면, 현재 수량과 별도로 앞으로의 대여 건수에 대한 전망이 필요함.')
bullet('', '예를 들어 현재 자전거 수가 같은 두 대여소라도, 다음 시간대의 예상 대여량이 다르면 점검 우선순위 검토에 다른 정보가 필요함. 이는 설명을 위한 가상 상황임.', sub=True)
sub('나. 예측 결과가 운영 판단에 연결되도록 서비스화할 필요')
bullet('활용 대상', '서울 전역 대여소를 확인하는 운영 담당자를 우선 사용자로 설정함. 시민용 대여 앱의 위치·수량 조회 기능을 대체하기보다 운영 검토를 지원하는 방향임.')
bullet('활용 장면', '날짜·시간을 선택하여 예상 이용량이 많은 대여소를 확인하고, 대여소별 과거 패턴과 예측 오차를 함께 비교함.')
bullet('기대 효과', '많은 대여소를 동일한 기준으로 살펴보고, 관심 지역과 시간대를 선별하는 데 필요한 정보를 제공함. 운영 효율 향상은 향후 현장 적용으로 확인할 사항임.')
sub('다. 과제의 목표를 측정 가능한 범위로 설정')
bullet('예측 대상', '대여소별·시간대별 실제 성사된 대여 건수로 한정함. 자전거 부족으로 이용하지 못한 잠재 수요나 재배치 최적화까지 포함하지 않음.')
bullet('검토 결론', '문제 정의 → 비교 실험 → 예측 지도·그래프 구현이 연결되는 학기 프로젝트로 추진 가능함. 공연·행사 정보는 필수 입력에서 제외함.')

section('2','기존 서비스·모델 대비 차별화 방향',True)
sub('가. 기존 기능의 역할과 본 과제의 추가 검토 사항')
table(['비교 대상','확인된 역할 또는 비교 기준','본 과제에서 다룰 내용'],[
    ('서울시 실시간 대여정보 [1]','현재 대여가능 수량·위치 등 상태 정보 제공','대여소별 향후 시간대의 대여 건수 전망 및 과거 예측 오차를 함께 제공'),
    ('UCI Seoul Bike 공개 예제 [5]','날씨·휴일을 포함한 시간별 대여량 자료. 실제 파일에는 대여소 ID 없음','서울시 대여소 ID·위치를 사용하여 대여소별 예측과 지도 화면으로 연결'),
    ('단순 과거 평균 기준(B0)','동일 대여소의 과거 요일·시간대 평균 등으로 예상값 산출','비선형 모델이 단순 반복 패턴보다 유리한지 먼저 확인'),
    ('개별 대여소 중심 기준(B1)','본 과제의 주 비교 모델. 해당 대여소 과거 패턴·달력·이용 가능한 날씨 입력','같은 모델에 주변·유사 대여소 특성을 더한 M1과 비교'),
],[3.7,6.65,7.05],9.2)
note('서비스 비교는 공식 공개 안내에서 확인한 기능을 기준으로 함. 운영기관 내부 모델이나 전체 앱 기능을 조사한 결과는 아니므로, “기존에 예측이 없다” 또는 “국내 최초”로 주장하지 않음.')
sub('나. 핵심 차별화는 새로운 알고리즘 명칭보다 검증할 질문에 둠')
bullet('개선 가설', '해당 대여소의 기록만 보는 것보다, 주변 대여소와 이용 패턴이 유사한 대여소의 과거 특성을 함께 보면 일부 시간·지역의 오차가 줄어들 수 있음.')
bullet('공간적 근거', '위치가 연결된 대여소 2,728곳 중 2,674곳에 500m 이내 다른 대여소가 있어, 주변 특성을 구성할 실질적인 연결 구조가 존재함.')
bullet('검증의 초점', '공간정보 활용 자체를 새로운 발명으로 제시하지 않고, 동일 조건에서 어느 대여소·시간대에 도움이 되는지와 개선이 없는 구간까지 분석함.')
sub('다. 서비스 차별화는 결과와 신뢰도를 함께 보여주는 데 둠')
bullet('예측의 설명', '지도에서 대여소를 선택하면 예상 대여 건수와 시간대별 패턴을 제시하고, 과거 날짜에서는 기본 모델·개선 모델·실제값을 함께 비교함.')
bullet('운영 판단의 범위', '예상 이용량이 많다는 이유만으로 자전거 부족이나 재배치 필요를 단정하지 않음. 수량·반납·이동을 함께 고려하는 의사결정은 후속 확장으로 구분함.')
box('제안하는 차별화 문장', '“서울 대여소별 대여량 예측에서 주변·유사 대여소 특성의 효과를 공정하게 비교하고, 그 결과를 운영 담당자가 확인할 수 있는 지도형 서비스로 구현한다.”')

section('3','데이터 확보 현황 및 사전 타당성 검증',True)
sub('가. 핵심 자료는 실제 원본을 확보하여 내용·연결 가능성을 확인')
table(['활용 자료','직접 확인한 내용','활용 목적'],[
    ('서울시 시간대별 이용정보 [2]\n2025년 1월 전체 CSV','1,610,875행 / 대여소 ID 2,730개\n이용건수 합계 1,704,190건\n날짜·시간·ID·건수 핵심 항목 결측 없음','대여소×시간별 정답 생성. 성별·연령·이용권으로 나뉜 행의 이용건수를 합산'),
    ('서울시 대여소 위치 [3]\n2024년 12월 XLSX','이용 ID 2,730개 중 2,728개 연결\n500m 내 이웃 수 중앙값 5곳\n미연결 ID: 4721·9980','과거 시점 지도·주변 대여소 거리 특성 구성'),
    ('기상청 ASOS 서울 108 [4]\n2025년 연간 CSV','8,760행·27항목. 1월 744시간 연결\n1월 기온·풍속·습도 결측 0\n압축파일 2단계 무결성 검사 통과','기상 입력 후보 확보. 사후 관측본이므로 예측 시점 이용 가능 여부 별도 반영'),
],[4.3,7.1,6.0],9.1)
sub('나. 사전 검증은 다음 순서로 수행')
bullet('원본 확인', '공식 제공 페이지에서 파일을 확보하고 압축 해제·인코딩·열 구조를 점검함. 이용정보의 “행 수”와 실제 “이용건수 합계”를 구분함.')
bullet('시간·공간 연결', '대여소 ID 표기를 통일하여 위치와 연결하고, 날짜·시간을 표준화하여 기상 시각과 결합함. 1월 이용자료에 등장한 서울 전체 744개 시각이 모두 날씨와 연결됨.')
bullet('검증 범위', '위 결과는 자료 활용 가능성에 대한 확인임. 모든 대여소의 모든 시간에 기록이 있다는 뜻이 아니며, 모델 정확도가 개선되었다는 증거도 아님.')
sub('다. 보조 자료는 효과와 공개 시점을 확인한 후 선택')
table(['자료','확인 결과 및 채택 판단'],[
    ('지하철 승하차 [6][7]','일별 자료: 2025년 1월 19,129행, D+3 공개 안내. 시간대별 자료: 621행의 월간 누계로 일별 시간 관측이 아님. 공개가 끝난 과거 패턴만 보조 입력으로 검토.'),
    ('서울 주요 공원 [8]','2026 하반기 목록 155개, 넓은 서울 주변 범위 내 유효 좌표 141개. 2025년 당시 상태가 보장되지 않아 최초 실험에서는 제외 가능.'),
    ('공휴일 [9]','공식 API 문서 확인, 정상 결과 목록은 미확보. 최신 2025년 공휴일·임시공휴일 결과 파일을 추가 확보하여 오픈소스 파생 달력과 대조 필요.'),
],[3.2,14.2],8.8)
note('따릉이 이용량은 현재 2025년 1월 전체 파일까지 확인함. 기상은 2025년 연간 자료를 확보함. 연중 성능 검증을 위한 여러 달의 이용량 확보는 다음 단계임.')

section('4','성과 검증 및 서비스 구현 방안',True)
sub('가. 성능 개선은 동일한 정답과 시험 조건에서 판단')
table(['비교 단계','모델·입력 구성','확인할 사항'],[
    ('B0\n단순 기준','공개 완료된 과거의 동일 요일·시간 평균 등','단순한 반복 패턴으로 설명되는 수준'),
    ('B1\n주 기준','LightGBM + 해당 대여소 과거 패턴 + 달력 + 이용 가능한 날씨','기본 예측 모델의 성능'),
    ('M1\n개선 후보','B1과 같은 LightGBM + 주변 거리·과거 패턴·유사 대여소 특성','추가 공간 특성의 순수한 개선 효과'),
],[2.6,9.3,5.5],9.2)
bullet('평가 지표', 'MAE를 주지표로 사용하여 시간당 평균 몇 건 틀렸는지 제시하고, 큰 오차에 민감한 RMSE를 함께 보고함.')
bullet('비교 조건', '학습 → 검증 → 시험을 시간순으로 분리하고 같은 시험 행을 사용함. B1/M1은 모델과 튜닝 예산을 같게 두며, 평균·군집 등은 학습기간에서만 산출함.')
bullet('결과 분석', '전체 평균 외에 출퇴근·고수요 시간 및 대여소별 오차와 표본 수를 확인함. 개선율 목표는 첫 비교 결과 이후 설정하며 불리한 구간도 함께 공개함.')
sub('나. 정답과 입력의 시간적 제약을 먼저 해결')
bullet('0건·누락 구분', '기록이 없는 시간이 실제 0건인지, 미운영·누락인지 확인하는 규칙이 필요함. 확인된 평가 범위를 표시하고 대여가 있었던 시간만 골라 평가하지 않음.')
bullet('공개 시점 준수', '공개 대여이력은 반기 갱신이므로, 학습 입력은 예측 당시 공개가 완료된 자료로 제한함. 실시간 직전 1시간 이용량을 전제로 한 비교를 하지 않음.')
bullet('기상 처리', '미래 실제 관측값을 입력하지 않고 이용 가능 시점에 맞춘 관측 또는 검증된 예보를 사용함. 강수 공란 690/744칸은 의미 확인 전 0으로 채우지 않음.')
sub('다. 첫 구현은 공개 자료 기반의 과거 재생·전망 대시보드로 한정')
table(['구현 항목','최소 기능'],[
    ('서울 지도·대여소 목록','날짜·시간 선택, 대여소별 예상 대여량 및 고이용 예상 대여소 확인'),
    ('대여소 상세·모델 비교','시간대별 그래프, 과거 실제값·B0/B1/M1 비교, 결측·적용 제외 표시'),
    ('데이터·예측 제공','공개 파일 집계 → 모델 학습 → 예측값 일괄 저장 → API → 화면 연결'),
],[4.1,13.3],9)

section('5','과제 평가 대응 및 회의 검토 사항',True)
sub('가. 평가 요소별 산출물')
table(['구분','평가 요소','준비할 근거·산출물'],[
    ('사전 40%','주제 적합성·창의성\n계획서·발표·질의응답','사용자·문제 정의, 기존 방식 대비 차이, 실제 데이터 확인 결과, 공정한 비교 계획 및 역할·일정'),
    ('최종 60%','모델 성능·Baseline 비교\n구현·적용·문서·발표·동료평가','동일 조건 비교표·오차 분석, 작동하는 지도와 그래프, GitHub/Wiki·재현 절차, 담당별 산출물과 기여 기록'),
],[2.7,5.8,8.9],9)
note('근거: 첨부 「2026_AI P-실무프로젝트_평가_260921.pdf」 2~3쪽. 새 딥러닝 알고리즘 개발을 반드시 요구하는 조건은 확인되지 않았으나, 기존 주제와의 차별성 설명은 필요함.')
sub('나. 오늘 회의에서 결정할 사항')
bullet('주제·범위', '대여소별 대여량 예측 및 운영지원 방향 채택 여부와, 주변·유사 대여소 특성 검증을 핵심 실험으로 둘지 결정함.')
bullet('데이터·일정', '이용량·위치·날씨·달력을 우선 사용하고, 보조 자료 포함 여부와 추가 월 확보·0건 판정 담당을 정함.')
bullet('역할 분담', '모델·평가 / 데이터 처리·API / 데이터 품질·출처 검토 / 화면·문서의 4개 역할별 담당과 산출물을 확정함.')
para('결정·보완 의견: _______________________________________________________\n담당 및 다음 점검일: ___________________________________________________',9.5,after=6)
sub('다. 제출 역산 일정')
para('10/19~21  주제·범위·담당 확정, 데이터 규칙·추가 확보 계획 정리\n10/22~24  계획서·발표자료·화면 초안 작성 및 검토\n10/25(일) 23:59  사전평가 제출: PPT·PDF·10분 이내 영상 URL TXT·계획서 WORD를 ZIP으로 팀장 1명 제출\n12/13(일) 23:59  최종 제출: 비교 실험·서비스 통합·문서화 완료 및 멘토 피드백 반영',9,after=5)
note('공휴일 보완: 공공데이터포털 [9]의 “공휴일 정보 조회”에서 solYear=2025, numOfRows=100, pageNo=1, 월은 비워 연간 조회. 1/27·6/3 추가 지정 반영 여부 확인. 인증키 대신 XML/JSON 또는 CSV 결과 파일만 공유.')
sub('참고 자료  ※ 제목 클릭 시 공식 출처로 이동')
refs=[
 ('[1] 서울시 실시간 대여정보','https://data.seoul.go.kr/dataList/OA-15493/A/1/datasetView.do'),
 ('[2] 시간대별 이용정보','https://data.seoul.go.kr/dataList/OA-15245/F/1/datasetView.do'),
 ('[3] 대여소 위치','https://data.seoul.go.kr/dataList/OA-13252/F/1/datasetView.do'),
 ('[4] 기상청 ASOS','https://data.kma.go.kr/data/grnd/selectAsosList.do?pgmNo=36'),
 ('[5] UCI Seoul Bike','https://archive.ics.uci.edu/dataset/560/seoul+bike+sharing+demand'),
 ('[6] 지하철 일별 승하차','https://data.seoul.go.kr/dataList/OA-12914/F/1/datasetView.do'),
 ('[7] 지하철 시간대별 월 누계','https://data.seoul.go.kr/dataList/OA-12252/S/1/datasetView.do'),
 ('[8] 서울 주요 공원','https://data.seoul.go.kr/dataList/OA-394/S/1/datasetView.do'),
 ('[9] 한국천문연구원 특일 정보','https://www.data.go.kr/data/15012690/openapi.do'),
]
for i in range(0,len(refs),3):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3)
    for j,(label,url) in enumerate(refs[i:i+3]):
        if j: run(p,'  |  ',8,color='666666')
        hyperlink(p,label,url)

props=doc.core_properties
props.title='서울시 공공자전거 대여량 예측 및 운영지원 서비스 추진 검토(안)'
props.subject='필요성·기존 방식 대비 차이·데이터 근거·성과 검증·회의 검토 사항'
props.author='P_Project'
props.keywords='따릉이, 추진 검토, 회의자료, 데이터 검증, 수요 예측'
props.created=props.modified=datetime(2026,10,19,tzinfo=timezone.utc)
path=OUT/'서울_따릉이_수요예측_추진검토안_20261019.docx'
doc.save(path)
print(path)
