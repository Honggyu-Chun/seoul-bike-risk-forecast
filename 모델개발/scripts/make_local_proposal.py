from pathlib import Path
from copy import deepcopy
from io import BytesIO
import hashlib,json,zipfile
from lxml import etree
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT.parent
TMP=ROOT/'reports/local_runs/proposal_authoring_20261020'
TMP.mkdir(parents=True,exist_ok=True)
REF=ROOT/'docs/reference_20261020/AI P-실무프로젝트 과제 계획서 양식 (1) (1).docx'
OUT=ROOT.parent/'과제제안서/이전작성본/따릉이_과제제안서_사전최종평가대응_20261020.docx'
OUT.parent.mkdir(parents=True,exist_ok=True)
result=json.loads((ROOT/'artifacts/improvement/final/results.json').read_text(encoding='utf-8'))
source=REF.read_bytes(); source_sha=hashlib.sha256(source).hexdigest()
with zipfile.ZipFile(BytesIO(source)) as z:
    inventory={n:{'bytes':len(z.read(n)),'sha256':hashlib.sha256(z.read(n)).hexdigest(),'policy':'editable' if n=='word/document.xml' else 'preserve'} for n in z.namelist()}
(TMP/'template_manifest.json').write_text(json.dumps({'path':str(REF),'sha256':source_sha,'parts':inventory},indent=2,ensure_ascii=False),encoding='utf-8')
doc=Document(BytesIO(source)); info=deepcopy(doc.tables[0]._tbl); section=deepcopy(doc.sections[0]._sectPr)
for el in list(doc.element.body):
    if el.tag!=qn('w:sectPr'):doc.element.body.remove(el)

def fmt(p,size=10.5,bold=False):
    p.paragraph_format.space_before=Pt(0)
    p.paragraph_format.space_after=Pt(6)
    p.paragraph_format.line_spacing=Pt(size*1.4)
    p.paragraph_format.widow_control=True
    pp=p._p.get_or_add_pPr()
    snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');pp.append(snap)
    pb=pp.find(qn('w:pBdr'))
    if pb is None:pb=OxmlElement('w:pBdr');pp.append(pb)
    for edge in ['top','left','bottom','right','between']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'nil');pb.append(e)
    for r in p.runs:
        r.font.name='맑은 고딕';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor(0,0,0)
        rp=r._element.get_or_add_rPr();fonts=rp.find(qn('w:rFonts'))
        if fonts is None:fonts=OxmlElement('w:rFonts');rp.insert(0,fonts)
        for name in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+name),'맑은 고딕')
        for name in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:fonts.attrib.pop(qn('w:'+name),None)
    return p
def p(text='',size=10.5,bold=False):return fmt(doc.add_paragraph(text),size,bold)
next_page=False
def h(text,level=1):
    global next_page
    a=doc.add_paragraph(text,style='Heading '+str(level));fmt(a,13 if level==1 else 11.5,True)
    if next_page:a.paragraph_format.page_break_before=True;next_page=False
    a.paragraph_format.space_before=Pt(12 if level==1 else 8);a.paragraph_format.space_after=Pt(7);a.paragraph_format.keep_with_next=True
    a._p.get_or_add_pPr().remove(a._p.pPr.find(qn('w:pBdr'))) if a._p.pPr.find(qn('w:pBdr')) is not None else None
    return a
def page():
    global next_page
    next_page=True
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    if widths:
        for row in t.rows:
            for c,w in zip(row.cells,widths):c.width=Inches(w)
    style_table(t)
    p('',size=2).paragraph_format.space_after=Pt(2)
    return t
def style_table(t,basic=False):
    props=t._tbl.tblPr
    borders=props.find(qn('w:tblBorders'))
    if borders is None:borders=OxmlElement('w:tblBorders');props.append(borders)
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b=borders.find(qn('w:'+edge))
        if b is None:b=OxmlElement('w:'+edge);borders.append(b)
        b.set(qn('w:val'),'single');b.set(qn('w:sz'),'5');b.set(qn('w:color'),'D9D9D9')
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        for old in list(trpr.findall(qn('w:trHeight'))):trpr.remove(old)
        cs=OxmlElement('w:cantSplit');trpr.append(cs)
        if i==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        seen=set()
        for j,c in enumerate(row.cells):
            if id(c._tc) in seen:continue
            seen.add(id(c._tc));c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr()
            for border in tcpr.findall('.//'+qn('w:tcBorders')+'/*'):
                border.set(qn('w:color'),'D9D9D9')
            shade=tcpr.find(qn('w:shd'))
            if shade is None:shade=OxmlElement('w:shd');tcpr.append(shade)
            shade.set(qn('w:fill'),'444444' if i==0 else ('F2F2F2' if (i%2==0 or basic and j==0) else 'FFFFFF'))
            shade.set(qn('w:val'),'clear')
            mar=tcpr.find(qn('w:tcMar'))
            if mar is None:mar=OxmlElement('w:tcMar');tcpr.append(mar)
            for edge in ['top','bottom','left','right']:
                elem=mar.find(qn('w:'+edge))
                if elem is None:elem=OxmlElement('w:'+edge);mar.append(elem)
                elem.set(qn('w:w'),'85');elem.set(qn('w:type'),'dxa')
            for a in c.paragraphs:
                fmt(a,10 if not basic else 9.5,i==0)
                a.paragraph_format.space_after=Pt(3);a.paragraph_format.line_spacing=Pt(13)
                if not basic and j>0 and len(a.text)<22:a.alignment=WD_ALIGN_PARAGRAPH.CENTER
                if i==0:
                    for r in a.runs:r.font.color.rgb=RGBColor(255,255,255)

title=doc.add_paragraph('따릉이 대여소 선택 지원 서비스 과제 계획서',style='Title');fmt(title,16,True)
p('AI P 실무프로젝트  |  따릉이 어디로  |  2026년 10월 20일',9.5)
doc.element.body.insert(len(doc.element.body)-1,info)
t=doc.tables[0]
t.cell(0,0).text='1. 기본 정보'
t.cell(1,1).text='따릉이 대여소 선택을 돕는 수량 부족 위험 예측 및 비교 서비스'
t.cell(2,1).text='Bike Availability Risk Forecasting and Station Comparison for Seoul Bike Riders'
t.cell(3,1).text='2026.9.1 ~ 2026.12.21'
t.cell(4,1).text=('팀 명:                          분반/조:                         \n'
'GitHub: https://github.com/Honggyu-Chun/P_Project\n'
'팀원 1  성명:             학번:             이메일:             \nGithub:                         역할안: 모델 및 성능 평가\n'
'팀원 2  성명:             학번:             이메일:             \nGithub:                         역할안: 데이터 및 백엔드\n'
'팀원 3  성명:             학번:             이메일:             \nGithub:                         역할안: 데이터 검증 및 문서\n'
'팀원 4  성명:             학번:             이메일:             \nGithub:                         역할안: 화면 및 사용자 검증')
t.cell(5,1).text='교수 멘토링 진행하지 않음';t.cell(5,3).text='현재 기업 멘토 없음\n산학 멘토링 대상 및 일정 확인 예정'
t.cell(6,1).text='[✓] 학생 자율 주제\n[  ] 기업 수요 제안 주제'
t.cell(6,3).text=''
t.cell(7,1).text='[  ] Physical AI\n[  ] AI 안전/보안/국방\n[✓] 행정서비스/공공 AI'
t.cell(7,3).text='[  ] 기업 제공 원천 데이터\n[✓] 공공 데이터/오픈 데이터셋\n[  ] 크롤링/자체생성 데이터'
t.cell(8,1).text='해당 없음'
style_table(t,True)
p('시민이 현재 수량과 다음 시간대의 부족 위험, 이동거리를 함께 비교하도록 돕고, 그 추가 가치와 한계를 같은 자료에서 검증한다.',10.5,True).paragraph_format.space_before=Pt(8)

page();h('2 프로젝트 개요 및 필요성')
h('2.1 과제 요약',2)
p('따릉이를 이용하려는 시민을 위해 주변 대여소의 현재 수량, 정확히 다음 한 시간 격자의 0대 위험, 출발지와의 직선거리를 비교하는 웹 서비스를 개발한다. 현재 수량 조회에 미래 위험과 선택 근거를 보태는 것이 목적이다. 서울시 공식 시간별 수량과 위치, 기상청 과거 관측을 활용한다. 고정한 2,729개 대여소에서 기본 모델과 개선 모델을 비교하고, 가상 요청으로 추천의 유효성과 추가거리를 함께 평가했다. 후속 시험에서 전체 F1은 72.959%에서 74.070%로 개선됐다. 과거 재생과 실시간 샘플 화면을 구현했으며, 남은 기간에는 사용성 확인과 운영 조건 검증, 재현 문서 및 발표 자료를 완성한다. 현장 대여 성공은 별도 검증 대상으로 둔다.')
h('2.2 개발 배경 및 문제 정의',2)
p('현재 자전거가 있다는 정보는 후보를 찾는 데 필요하다. 그러나 시민이 실제로 선택해야 하는 것은 “조금 더 가까운 곳”과 “다음 시간대에도 자전거가 남을 것으로 예상되는 곳” 사이의 대안이다. 현재 수량 하나만으로 이 선택을 충분히 설명할 수 있는지 확인할 필요가 있다. 공식 실시간 자료는 현재 수량과 위치를 제공하며, 본 과제는 그 자료를 시민의 선택 지원으로 연결한다. 기존 앱 전체 기능이나 운영기관 내부 모델의 부재를 전제하지 않는다. [1][2][3]')
table(['이용 장면','시민이 판단할 내용','제공할 정보'],[
['다음 시간대 이용을 미리 준비','가까운 후보 중 어디를 우선 확인할지','현재 수량과 한 시간 격자의 위험 신호'],
['가까운 곳의 수량이 적음','추가 이동을 감수할 만한 대안이 있는지','주변 후보의 수량과 추가 직선거리'],
['정보가 오래됐거나 연결 실패','추천을 믿고 이동할 수 있는 상태인지','수신시각과 경과시간 및 재확인 안내']], [1.35,2.25,2.9])
p('필요성은 과거 자료에서도 확인할 수 있다. 새 시험의 현재 양수 기록 1,776,483건 중 51,237건은 다음 기록에서 0대가 됐다. 현재 1~3대인 기록에서는 이 비율이 약 12.35%였다. 가까운 현재 양수 대여소를 고르는 기준도 공통 가상 요청 10,443건 중 618건에서 다음 기록에 수량이 없었다. 현재 수량만으로 미래 상태가 보장되지 않는다는 근거다. 다만 실제 시민의 불편 빈도나 대여 실패율을 뜻하지 않는다. [9]')
p('따라서 본 과제의 가치는 예측 점수 자체보다 “현재 수량에 더해 무엇을 보여 주면 더 합리적인 선택을 도울 수 있는가”를 검증하는 데 있다. 실패 사례도 공개하고, 추가 이동을 늘려서 얻은 효과인지 구분한다. 사회적 가치는 불필요한 이동을 줄일 가능성이며, 경제적 효과나 실제 시간 절감은 현장 자료를 확보한 후 산정한다.')

page();h('3 프로젝트 목표 및 독창성')
h('3.1 최종 개발 목표',2)
table(['목표','완료 기준','2026년 10월 20일 기준'],[
['선택 지원 화면 완성','검색 → 후보 비교 → 상세 근거 → 선택까지 실행','과거 재생과 오프라인 시연 구현'],
['AI의 추가 가치 검증','같은 시험의 B0·B1·B1+·M2 및 R0·R1u·Rq·R2 비교','고정된 후속 시험 결과 확보'],
['실사용으로 이어질 조건 확인','소규모 사용성 기록, 정보 신선도와 오류 대응 점검','화면 구현, 사용자·전역 운영 검증은 계획'],
['팀이 이어갈 수 있는 결과물','실행 안내·출처·시험 규칙·개인 기여·발표 증빙 정리','복원 자료를 기반으로 보완']], [1.25,2.65,2.6])
p('초기 팀 연구 목표는 전체 F1 +5%p, 재현율 비악화, R0 대비 전체 유효추천율 +5%p, 평균 추가 직선거리 100m 이하이다. +5%p는 75%에서 80%처럼 비율의 차이를 뜻한다. 이 수치는 학교가 정한 합격선이 아니며, 실제 달성 여부는 6.2와 6.3에서 그대로 판정한다. 이미 확인한 결과에 맞춰 목표를 낮춰 달성으로 바꾸지 않는다.')
h('3.2 독창성 및 선행 방식과의 차별성',2)
table(['비교 대상','활용하는 점','본 과제에서 더 확인할 점'],[
['공식 현재 수량 조회 [3]','수량과 위치로 가까운 후보 파악','다음 격자 위험과 추가거리의 선택 지원'],
['자기 정보 기반 AI','대여소의 최근 수량과 시간 패턴','주변·전환·과거 날씨 및 모델 결합의 실제 개선 폭'],
['수량과 거리만 쓰는 규칙 Rq','AI 없이도 가능한 강한 비교 기준','비슷한 유효율에서 관측된 거리 차이와 불확실성'],
['교통 시공간 예측 연구 [7][8]','주변 흐름과 시간 이력의 활용 아이디어','교통량 회귀를 다음 기록 0대 분류로 그대로 동일시하지 않음']], [1.4,2.15,2.95])
p('자체 기여는 새로운 범용 알고리즘의 발명보다, 시민의 선택 문제에 맞춘 정답 생성, 누락 처리, 시간순 검증, 주변 정보 구성, 위험·거리 추천 평가와 근거 화면의 결합이다. LightGBM·CatBoost 등 공개 구현을 활용하며, 외부 논문의 성능을 이 데이터에서 재현한 성능으로 제시하지 않는다.')
p('주변 정보를 더한 기존 M1만으로는 충분한 개선을 확인하지 못했다. 이에 단순 규칙과 더 조정한 기본 AI까지 대조군을 확장했다. 어떤 기술이든 추가 효과가 작거나 없으면 그 결과를 공개하고, 운영 복잡도에 비해 필요한지 판단한다.')

page();h('4 세부 개발 내용 및 방법')
h('4.1 시스템 아키텍처 및 파이프라인',2)
p('공식 원본 수집 → 파일 해시 및 관측 검수 → 대여소 번호와 시간 격자 연결 → 현재까지의 특징 생성 → 학습 및 검증 → 모델과 정책 고정 → 이후 기간 평가 → API와 시민용 화면',10.5,True)
p('학습 단계와 서비스 조회 단계를 분리한다. 과거 화면은 저장된 시험 예측과 실제 수량을 읽으므로 시연할 때 재학습할 필요가 없다. 정답은 검증 스위치를 켰을 때 구분해서 보여 준다. 서버는 입력 시각까지의 실제 이력만 반환하며 미래 수량을 설명 입력에 섞지 않는다.')
table(['경로','화면에 제공하는 내용','해석 범위'],[
['기존 3월 M1 재생','대여소 검색·후보·24시간 이력·500m 이웃·사례','저장된 과거 예측의 연구 시연'],
['후속 4~5월 M2 결과','새 기간 모델·추천 비교와 보고서','별도 시험의 수치이며 현재 지도 예측과 구분'],
['실시간 B1S 샘플','공식 샘플 5곳 현재 수량과 기본 모형 실행','실시간 M1·M2 정확도 검증과 구분']], [1.3,2.7,2.5])
h('4.2 적용 기술 스택 및 예외 처리 방안',2)
table(['구성','기술 및 역할'],[
['프론트엔드','HTML·CSS·JavaScript, 반응형 화면과 SVG 참고 지도, 검색·즐겨찾기·CSV'],
['백엔드','Python 3.12, FastAPI·Uvicorn, 조회와 입력 검증 및 상태 반환'],
['데이터 저장','pandas·NumPy·Parquet·NPZ·JSON 파일, 출처·스냅샷·시험 산출물 분리'],
['AI와 검증','LightGBM·CatBoost, scikit-learn 지표·확률 보정, pytest·Playwright'],
['실행 및 배포','uv.lock의 고정 의존성, 로컬 서버와 독립형 HTML. 공개 운영은 후속 조건 확인 후 결정']], [1.15,5.35])
p('예외 상황은 사용자가 선택을 바꿀 수 있도록 처리한다. 후보가 없으면 추천을 보류하고, 누락 수량은 0으로 표시하지 않는다. 과거 조회 오류에는 직전 성공 조건을 함께 유지한다. 실시간 갱신 실패 시 남은 값이 이전 수신 자료임을 표시하고, 수신 후 120초가 지나면 재조회를 안내한다. 자동 갱신은 해당 화면을 보는 동안에만 수행한다.')
p('위치도와 500m 범위는 직선거리 기준이며 외부 지도 연결로 실제 길을 확인하도록 한다. 주변 이력은 참고 근거이고, 개별 예측의 원인을 입증한 설명으로 표시하지 않는다.')

page();h('4.3 데이터 수집 및 학습과 전처리',1)
table(['공식 데이터','확보 및 확인 범위','분석에서의 역할'],[
['시간별 수량 OA-22382 [1]','2025년 1분기 5,835,300행, 원본 번호 2,767개. 후속 2분기 ZIP 추가','시간 이력과 정확히 다음 격자의 0대 정답'],
['대여소 위치 OA-13252 [2]','2024년 12월 위치자료를 수량 번호와 연결','후보 검색, 500m 이웃과 직선거리'],
['기상청 서울 ASOS 108 [4]','2025년 8,760시간. 관측 1·24시간 전의 11개 특징','후속 과거 날씨 후보 비교. 게시 지연은 별도 확인'],
['실시간 bikeList [3]','공식 공개 샘플 5곳과 별도 B1S 실행 기록','현재 입력 연결과 신선도 상태 시연']], [1.35,2.7,2.45])
h('대상과 정답을 만드는 규칙',2)
p('대상은 1~2월 학습자료에서 336시간 이상 관측, 24시간 이상 양수, 수량 변동, 유효 위치 조건을 충족한 2,729개 대여소로 고정했다. 제외된 38곳은 자료 조건에 따른 것이며 공식 운영 여부 판정이 아니다. 시험 점수를 보고 어려운 대여소를 빼지 않는다. [9]')
p('입력 시각 t의 같은 대여소에 대해 정확히 t+1시간의 실제 기록이 있을 때만 정답을 만든다. 다음 기록이 0대면 양성, 1대 이상이면 음성이다. 1분기 도시 공통 누락 50시간은 0대로 채우지 않는다. 수량·위치·원본 파일 해시와 대상 선정표를 함께 보관한다.')
p('모든 이동평균·변화량·주변값은 현재 또는 과거 관측으로만 만든다. 전환 이력은 이미 관측한 q(t-1)에서 q(t)로의 변화까지만 누적한다. 사후 관측을 이용해 미래 값을 입력하는 누수를 막는다. 비와 눈의 미확인 값도 임의로 0으로 바꾸지 않는다.')
h('자료 해석과 보안',2)
p('공식 설명은 대여가능 수량이지만 원본 열 이름은 거치대수량이다. 시간 안의 집계 시점과 운영 상태 의미는 제공기관에 추가 확인할 항목이다. 정답을 실제 대여 실패·고장·한 시간 내내 품절과 동일시하지 않는다. 기상 아카이브의 사후 수정과 당시 제공본의 동일성, 게시 지연도 아직 확인하지 못했다.')
p('대여소 단위 공개 자료를 사용하고 실제 개인 이동기록은 수집하지 않는다. 위치 권한은 사용자가 선택하며 즐겨찾기는 브라우저에 저장한다. 공유 링크에 위치가 포함됨을 안내한다. 발급키는 서버 환경변수로만 읽고 프론트엔드·공개 저장소·보고서에 노출하지 않는다. 사용자 조사에는 동의를 받고 익명 결과만 정리한다.')

page();h('5 산학 멘토링 및 추진 일정')
h('5.1 주차별 추진 일정',2)
p('수행 기간은 2026.9.1부터 2026.12.21까지 유지한다. 현재 확보한 코드와 실험 결과를 출발점으로 삼고, 남은 기간은 사용성·운영 조건·발표와 재현 자료를 완성하는 데 사용한다. 과거 주차의 활동이나 개인 기여를 소급해 완료 기록으로 만들지 않는다. [10][11]')
table(['주차와 시점','수행 내용','완료 확인 자료'],[
['1~8주 및 현재','기획·자료 확보·시제품·후속 비교 결과를 정리하고 로컬로 인수','출처·코드·기존 결과·로컬 검증 기록'],
['9주 사전평가\n10/25 제출','계획서·PPT·영상 정리, 문제와 비교 설계 및 현재 결과 설명','PPT·PDF·10분 이내 영상 URL TXT·계획서 WORD'],
['10주 구현 정리','검색→비교→상세→사례→결과의 전체 흐름과 예외 상황 점검','기능 체크리스트와 브라우저 시연 기록'],
['11주 1차 산학 멘토링','시민 이용 상황, 시간 격자 한계, 대조군과 운영 조건 검토','질문·의견·채택 여부·담당·반영 기한'],
['12~13주 및 2차 멘토링','소규모 사용성 확인, 현장 검증 가능성 검토, 반영 결과 재점검','익명 관찰 기록·이슈·반영 전후 증거'],
['14~15주\n12/13 제출','오류 수정·실행 안내·공개 가능 파일 정리·최종 자료 및 영상 제작','최종보고서·PPT·PDF·15분 이내 영상 TXT·멘토 정보 반영 계획서'],
['16주\n12/15~12/21','분반 공지에 맞춰 7분 내외 발표와 질의응답','최종 시연·발표·동료평가 및 인수 자료']], [1.25,2.65,2.6])
p('사전 발표는 10/27~11/2, 최종 발표는 12/15~12/21 범위에서 분반별로 다르다. 사전 제출은 10/25 일요일 23:59, 최종 제출은 12/13 일요일 23:59에 팀장 1명이 ZIP으로 제출한다. 정확한 분반별 발표 일시와 파일명은 사이버캠퍼스 공지를 확인한다. [10]')
p('10주차 이전에 문제 정의·구현·실험이 일정 수준 준비돼야 한다는 OT 취지에 맞춰, 이미 있는 과거 시제품과 고정 실험 결과를 멘토링 자료로 제시한다. “구현 80%” 같은 양식 예시 수치를 실제 진척도로 기입하지 않는다.')

page();h('5.2 멘토링 준비 및 피드백 환류',1)
table(['시점','사전 질의','피드백을 반영할 결과'],[
['11주 1차','한 시간 격자 예측이 유용한 시민의 준비 상황은 무엇인가. 실제 대여 실패를 어떻게 수집할 것인가','서비스 이용 장면·표시 문구·검증 범위 확정'],
['11주 1차','R0·Rq·R1u 대조군과 현재 양수 전환 지표가 충분한가','비교 설계·한계·오류 분석 검토 기록'],
['13주 2차','소규모 사용성 결과와 정보 신선도 표시가 충분한가. 전역 운영의 선행 조건은 무엇인가','화면 수정·운영 조건·최종 주장 범위 확정']], [1.0,3.1,2.4])
p('멘토링 후 48시간 이내에 날짜·멘토 정보·질문·답변·반영 여부·담당·기한을 회의록으로 정리한다. 채택한 의견은 GitHub 이슈와 변경 내역에 연결하고, 보류한 의견도 이유를 기록한다. 2차 멘토링에서는 1차 의견의 반영 전후를 보여 준다. 실제 실시 후에만 기업 멘토 정보와 증빙을 최종 계획서에 넣는다.')
p('교수 멘토링은 선택 사항이며 본 팀은 진행하지 않는다. 현재 기업 멘토는 없으므로 산학 멘토링 대상과 일정을 학교에 확인한다. OT에 명시된 11주와 13주 2회 멘토링은 준비 계획에 유지한다. 실제 실시한 뒤에만 성명·소속·의견을 기록하며, 비대면 증빙은 학교 양식이 허용하는 전자서명·화면·피드백 기록 중 확보한 자료를 사용한다. [10][11][12]')
h('5.3 4인 역할 분담 및 협업 계획',2)
table(['역할안','책임 결과물','교차 검토'],[
['A 모델 및 평가','예측 코드·실험 요약·오류 분석·시험 고정 기록','C가 원본 결과와 수치를 대조'],
['B 데이터 및 백엔드','수집·연결·조회 API·Windows 실행 안내','A가 시간 정렬, D가 화면 연결 확인'],
['C 데이터 검증 및 문서','검수 스크립트·누락/중복/위치 점검·출처·평가표','B가 자료 처리 규칙을 검토'],
['D 화면 및 사용자 검증','작동 화면·사용성 과제·익명 관찰·시연 시나리오','C가 표시 수치와 한계를 확인']], [1.3,3.0,2.2])
p('이는 제안된 분담이며 실제 담당자는 팀에서 확정한다. 기존 코드와 AI 지원 산출물의 출처를 구분하고, 각자의 이슈·수정·검토·시연 기록으로 실제 기여를 남긴다. 매주 결과물 하나를 서로 실행해 보고 담당 외 영역도 설명할 수 있도록 한다. 계획된 역할을 실제 참여 실적으로 바꾸어 동료평가에 사용하지 않는다.')

page();h('6 기대효과와 정량적 성능 검증')
h('6.1 벤치마크 및 객관적 평가 설계',2)
p('모델의 정답은 다음 기록 0대 여부인 이진 분류다. F1은 위험 알림의 정확성과 실제 위험을 찾는 정도의 균형이며, 실제 대여 성공률이 아니다. 위험 사례가 새 시험의 약 9.399%이므로 단순 정확도만으로 평가하지 않는다.')
table(['지표','계산 및 해석','보고할 내용'],[
['F1','2 × TP ÷ (2 × TP + FP + FN)','주지표 및 같은 시험의 개선 %p'],
['정밀도와 재현율','TP/(TP+FP), TP/(TP+FN)','오경보 FP와 놓친 위험 FN을 함께 제시'],
['확률 보조 지표','평균정밀도와 Brier 확률 오차','순위와 확률의 신뢰도를 구분'],
['유효추천율','다음 기록 양수인 곳 추천 건수 / 전체 요청','추천 보류를 분모에 남김'],
['추가 직선거리','R0와 같은 요청에서 추천 거리 차이의 평균','양쪽 모두 추천한 요청 수·보류율 병기']], [1.2,3.1,2.2])
p('TP는 실제 0대를 경고한 수, FP는 실제 양수인데 경고한 수, FN은 실제 0대를 놓친 수이다. 지표 분모가 0이면 0으로 처리한다. 수량 회귀의 MAE·RMSE·L2와 분류 F1을 같은 지표로 바꾸어 설명하지 않는다.')
table(['단계','기간과 용도'],[
['1차 탐색','1~2월 학습, 3/1~3/14 검증. 기존 3/15~3/31 시험은 이미 확인한 자료'],
['후속 학습과 선택','1~3월 학습. 4/1~4/14에서 모델·혼합비·경고 기준 선택'],
['확률 보정과 정책','4/1~4/7 확률 보정, 4/8~4/14 추천의 위험·거리 가중치 선택'],
['고정 후 최종 시험','2025/4/15~5/14, 1,960,441건. 실제 0대 184,264건']], [1.55,4.95])
p('모델·임계값·정책·보정기·코드와 설정 해시를 저장한 뒤 시험을 한 번 평가했다. 현재는 그 시험도 열린 자료다. 추가 선택에는 별도의 미사용 기간과 새 프로토콜이 필요하다. 5/15 이후 원본이 있다는 사실만으로 모든 후속 과정에서 미사용임이 증명되지는 않으므로 접근 이력을 먼저 점검한다. [9]')

page();h('6.2 모델 성능과 목표 판정')
p('B0는 현재 상태 유지, B1은 자기 대여소 정보 16개의 LightGBM, B1+는 같은 자기 정보를 더 조정한 대조군이다. M2는 시간·전환·주변·과거 날씨 LightGBM과 CatBoost를 50%씩 결합했다. M2 경고 임계값은 검증에서 고정한 0.39다. 1차 후보 45개와 후속 후보 15개를 실제 비교했으며, 모든 후보의 탐색·표본·계산 예산이 같았다는 뜻은 아니다.')
m=result['metrics']
table(['같은 새 시험','F1','정밀도','재현율','놓친 0대'],[[name]+[f'{m[key][k]*100:.3f}%' for k in ['f1','precision','recall']]+[f"{m[key]['fn']:,}"] for name,key in [('B0 현재 유지','B0'),('B1 기본 AI','B1'),('B1+ 조정 AI','B1_tuned'),('M2 제안','M2')]], [1.45,1.0,1.0,1.0,1.05])
p('M2는 B1 대비 F1 +1.1109%p, 재현율 +1.5049%p였다. 놓친 위험은 50,960건에서 48,187건으로 2,773건 줄고, 오경보도 47,852건에서 47,086건으로 766건 줄었다. 날짜 30개를 묶은 2,000회 짝지은 재표집의 탐색적 95% F1 차이 범위는 +0.7543~+1.4236%p였다. 해당 계절·코호트에서의 개선 증거이며, 연중·현장의 보장은 아니다. [9]')
h('시민이 고를 수 있는 현재 양수 구간',2)
s=result['subgroups']['current_positive']
table(['현재 양수 구간','F1','전환 재현율','놓친 전환'],[[name,f"{s[key]['f1']*100:.3f}%",f"{s[key]['recall']*100:.3f}%",f"{s[key]['fn']:,}"] for name,key in [('B1','B1'),('B1+','B1_tuned'),('M2','M2')]], [2.1,1.2,1.6,1.6])
p('현재 양수 구간 1,776,483건의 실제 0대 전환은 51,237건이다. M2는 개선됐지만 전환의 약 85.2%를 여전히 놓쳤다. 따라서 “위험 낮음”을 대여 보장으로 표시하지 않는다. 이 구간의 F1 +10.647%p를 전체 F1 +5%p 목표 달성으로 대체하지 않는다.')
table(['초기 팀 목표','같은 새 시험의 결과','판정'],[
['전체 F1 +5%p','+1.1109%p','미달'],['재현율 비악화','+1.5049%p','충족'],['전체 유효추천율 +5%p','+4.0218%p','미달'],['전체 평균 추가거리 100m 이하','32.18m','충족']], [3.0,2.3,1.2])
p('기존 3월 M1의 F1은 B1 75.741% → M1 75.766%로 +0.0245%p였고, 확실한 개선을 입증하지 못했다. 3월 점수와 새 4~5월 M2 점수를 직접 빼서 모델 개선 폭으로 계산하지 않는다.')

page();h('6.3 추천 효과와 거리의 함께 검증')
p('고정 가상 출발지 120개와 평일 7·8·18·19시, 500m 안의 현재 양수 후보를 사용했다. 예정 요청 10,560건 중 공통 후보가 없는 117건을 제외한 10,443건으로 모든 방법을 비교했다. 추천 보류는 이 전체 분모에 남긴다. 후보는 다음 정답 기록이 있는 과거 자료로 제한되므로 미래 정답 누락이 있는 실제 운영 상황 전체의 성능은 아니다. [9]')
rq=result['recommendation']['metrics']
table(['방식','유효 건수 / 전체','유효추천율','추가거리','보류율'],[[key,f"{v['valid_recommendations']:,} / {v['requests']:,}",f"{v['effective_rate']*100:.3f}%",f"{v['extra_distance_m']:.2f}m",f"{v['abstention_rate']*100:.3f}%"] for key,v in rq.items()], [0.6,1.95,1.45,1.25,1.25])
p('R0는 현재 양수 중 최근접, R1은 B1 위험 제외 후 최근접이다. R1u는 B1에도 R2와 같은 보정·위험/거리 효용 규칙을 적용한다. Rq는 AI 없이 현재 수량과 거리만 사용하며, R2는 M2의 보정 위험과 거리를 사용한다. R1u로 추천 규칙 변화와 모델 변화의 효과를 분리해서 비교한다.')
p('R2는 R0보다 유효한 추천이 420건 늘어 +4.0218%p였다. 같은 정책 계열의 R1u보다 40건, +0.3830%p 늘었다. 날짜별 짝지은 2,000회 재표집의 탐색적 95% 범위는 각각 +3.441~+4.543%p, +0.210~+0.576%p다. R0 대비 차이 전체를 모델 자체의 효과라고 설명하지 않는다.')
p('강한 단순 기준 Rq는 R2보다 유효 추천이 5건 많았다. R2의 평균 추가 직선거리는 32.18m로 Rq의 75.45m보다 관측상 약 43.27m 짧았다. R2−Rq 유효추천율 차이의 범위는 -0.239~+0.144%p로 0을 포함한다. 더 높은 유효율, 통계적 동등성 또는 실제 보행거리 절감을 입증한 것으로 표현하지 않는다.')
h('도움이 클 수 있는 조건과 남는 한계',2)
p('최근접 후보가 현재 1~3대인 2,677건에서는 R0 81.584%, R2 95.219%로 +13.635%p였다. 그러나 이 구간의 R2 평균 추가거리는 103.87m로 100m를 넘는다. 전체 지표의 목표 판정과 구분하고 다음 사용자 검증의 가설로 활용한다.')
p('기존 3월 추천 시험은 R0 95.176%로 +5%p 자체가 100%를 넘고, 미래 정답을 쓰는 참고 상한에서도 최대 +4.039%p였다. 새 4~5월 시험의 참고 상한 여유는 +5.315%p이므로 +5%p는 상한 안에 있지만 이번 결과는 미달이다. 이 상한은 실제 운영에서 사용할 예측기가 아니다.')

page();h('6.4 기대효과 및 결과물 활용 계획')
p('시민에게는 “현재 몇 대인가”와 “어떤 대안을 더 확인할 것인가”를 같은 화면에서 판단할 수 있는 기능을 제공한다. 공공서비스 관점에서는 공개 데이터를 실제 선택 문제로 연결하고, 추가 이동과 품절 위험 사이의 기준을 검토할 수 있다. 학습 측면에서는 자료 검수부터 모델·화면·테스트·문서까지 이어지는 AI 개발 과정을 4명이 나누어 경험한다. 이용 확대나 비용 절감의 규모는 아직 추정하지 않는다.')
h('사용자 검증과 실서비스 전환 조건',2)
p('소규모 탐색은 따릉이 이용 경험이 있는 자발적 참여자 3~5명에게 시행할 계획이다. 같은 과거 상황을 현재 수량·거리 화면과 위험을 추가한 화면으로 번갈아 제시한다. 후보 선택 이유, 과제 완료 시간, 표시 의미 이해, 허용 추가 이동을 기록하고 제시 순서를 교차한다. 대여 보장·도보시간·위험 신호를 오해한 경우를 오류로 분류해 문구를 수정한다. 이는 사용성 탐색이며 인구 전체의 효과나 대여 성공률 검증이 아니다.')
p('전역 서비스로 확대하려면 발급키 전체 조회와 장기 호출 조건, 대여소 연결, 원천 갱신 지연, 과거 이력 누적을 확인해야 한다. 한 번의 조회로 1·24·168시간 이력을 대신하지 않는다. 5~10분 예측은 그 간격의 기록과 새 정답, 미사용 미래 기간이 확보된 뒤 별도 목표로 검증한다. M2에 쓰인 날씨도 당시 제공 가능성과 게시 지연을 확인해야 한다.')
h('6.5 사전발표 평가항목 충족 계획',2)
table(['공식 항목','이 계획서의 대응과 발표 증빙'],[
['주제 적합성','2장과 4장: 시민의 선택 문제·공식 자료·다음 기록 정답·한 학기 범위를 연결. 실제 입력→비교 화면 시연'],
['창의성','3장과 6.3: 현재 수량 조회에 미래 위험·거리 선택을 결합. Rq·R1u 대조군으로 자체 기여와 추가 가치 설명'],
['발표 자료 구성 및 Q&A 응답','문제→사례→자료→방법→동일 시험 결과→한계→실행계획 순서. 아래 질의응답 근거와 발표 리허설'],
['프로젝트 계획서 작성','학교 양식 1~7항, 고정 수행 기간·정량 목표·WBS·멘토링·역할·출처를 완성하고 개인정보는 팀이 확정하여 입력']], [1.55,4.95])
p('사전평가 반영비율은 총 40%이며 항목별 세부 배점이나 합격선은 제공 자료에 없으므로 임의로 만들지 않는다. 사전 발표는 5분 내외, 제출 영상은 10분 이내다. 계획과 이미 측정한 결과를 화면·표에서 구분한다. [10]')

page();h('6.6 최종발표 평가항목 충족 계획')
table(['공식 항목','제시할 증거와 완료 방법'],[
['모델 성능','6.2의 동일 시험 F1·정밀도·재현율·FN/FP, 현재 양수 구간 및 목표 달성표. results.json과 일별 결과로 검산'],
['벤치마크 대비 우수성 및 객관적 성능 검증','B0·B1·B1+·M2와 R0·R1u·Rq·R2의 같은 기간 비교. 시간 분리·선택 해시·날짜 재표집. 향상·미달·불리한 비교를 함께 공개'],
['구현 완성도','검색→범위→후보 비교→상세 이력→사례→결과·보고서 다운로드를 끝까지 시연. API 입력·누락·오류·모바일·오프라인 검증 기록'],
['적용 가능성','6.3의 유효율·보류·거리와 6.4의 사용성 기록. 현재 표본 범위·직선거리·실시간 전환 조건을 명시하고 현장 효과는 별도 검증'],
['문서화 및 오픈소스화','README·설치 명령·API·데이터 출처·모델 설명·시험 규칙·한계·라이선스·Wiki 준비. 실제 공개 범위와 커밋·릴리스·Wiki URL을 최종 자료에서 확인'],
['발표 자료 구성 및 Q&A 응답','새 M2 수치와 기존 M1 지도 및 B1S 샘플을 구분. 성공·실패 사례, 목표 미달, 강한 단순 기준을 포함. 7분 내외 발표·15분 이내 영상 리허설'],
['동료 평가','5.3의 책임 결과물을 실제 이슈·커밋·리뷰·회의록과 연결. 각자의 구현·검토·발표 기여를 확인하고 계획과 실적을 구분']], [1.7,4.8])
p('최종평가 반영비율은 총 60%이다. Baseline 대비 개선 근거는 일부 확보했지만 모든 기준에 대한 평가 결과가 확정된 것은 아니다. 특히 현장 적용성과 실제 팀 기여, 공개 Wiki 및 발표 자료는 남은 일정에서 증거를 완성한다. [10][11]')
h('문서화와 공개를 완료하는 방법',2)
p('코드·원본 데이터·학습 모델·시험 결과의 공개 범위를 나눈다. 소스에는 사용 라이브러리와 고지문을 정리하고 팀 작성 코드의 라이선스를 확정한다. 데이터는 공식 이용조건과 출처를 기록하고 재배포 가능 범위를 확인한다. 공개할 수 없는 파일은 다운로드·해시 검증·복원 명령으로 재현 경로를 제공한다.')
p('현재 GitHub 이전 브랜치는 전체 백업을 전달하는 구조다. 최종 오픈소스화에서는 앱 소스와 실행 안내를 쉽게 찾을 수 있는 구조, 사용 예시, 환경·데이터 조건, 고정 결과표를 정리한다. OT가 요구한 Wiki에는 제안서·일정·결정·진행 기록을 관리하고 발표 때 URL을 공유한다. 이미 공개·배포를 완료한 것으로 기재하지 않는다.')

page();h('6.7 발표 구성 및 질의응답 준비')
p('사전 5분 발표안은 문제와 시민 장면 1분 → 자료·정답·차별성 1분 → 시제품 1분 → 현재 비교 결과와 한계 1분 → 남은 일정·멘토링·역할 1분으로 구성한다. 최종 7분은 문제·목표 1분 → 비교 설계 1분 → 전체 흐름 시연 2분 → 모델·추천 결과 2분 → 실패·활용·기여 1분으로 준비한다.')
table(['예상 질문','답변의 핵심과 근거'],[
['현재 수량만 봐도 되지 않는가','현재 양수→다음 0대 전환과 R0의 과거 오류가 있어 미래 비교를 시험한다. 실제 불편 빈도는 사용자 조사로 확인한다. 2.2·6.3'],
['AI를 쓰는 이유가 충분한가','R2는 R0보다 유효율이 높았지만 Rq보다 높지는 않았다. 같은 규칙의 R1u 비교와 거리 차이를 함께 보고한다. 6.3'],
['주변 정보만으로 좋아졌는가','M2는 학습기·시간 특징·전환·날씨·탐색도 달라졌다. 전체 향상을 주변 정보의 단독 효과로 주장하지 않는다. 3.2·6.2'],
['한 시간 예측이 도착 시점과 같은가','정확히 다음 격자 기록의 위험이다. 준비 단계 선택 보조이며 5~10분 정확도는 별도 자료로 검증해야 한다. 2.1·6.4'],
['왜 F1을 쓰고 정확도를 주지표로 쓰지 않는가','0대가 약 9.4%로 적어 모두 양수라 해도 정확도가 높을 수 있다. F1과 놓침·오경보를 함께 제시한다. 6.1'],
['목표 +5%p를 달성했는가','두 전체 +5%p는 미달, 재현율 비악화와 전체 평균거리 조건은 충족했다. 하위 구간으로 목표를 대체하지 않는다. 6.2'],
['시험을 다시 써서 점수를 올린 것 아닌가','후속 학습에 이미 본 3월을 사용한 사실을 밝히고, 4/15~5/14 전에 고정한 선택·모델·설정 해시를 제시한다. 6.1·[9]'],
['실시간 서울 전체가 작동하는가','이번 범위는 과거 M1 재생·M2 결과표·B1S 공개 샘플이다. 전역·현장·짧은 간격 성능은 별도 검증이다. 4.1·6.4']], [1.8,4.7])
p('시연에는 도움 사례와 불리한 사례를 모두 포함한다. 네트워크 장애 시 독립형 HTML의 실제 과거 자료로 설명한다. 영상·PPT·PDF·계획서의 수치는 같은 results.json에서 확인하고, 모델 이름과 시험기간을 표 제목에 항상 붙인다.')

page();h('7 참고문헌 및 검증 자료')
refs=[
'[1] 서울 열린데이터광장. 서울시 대여소별 공공자전거 따릉이 대여가능 수량 OA-22382. 2025년 1·2분기 원본. https://data.seoul.go.kr/dataList/OA-22382/F/1/datasetView.do',
'[2] 서울 열린데이터광장. 서울시 공공자전거 따릉이 대여소 정보 OA-13252. 사용본 2024년 12월. https://data.seoul.go.kr/dataList/OA-13252/F/1/datasetView.do',
'[3] 서울 열린데이터광장. 서울시 공공자전거 따릉이 실시간 대여정보 OA-15493. https://data.seoul.go.kr/dataList/OA-15493/A/1/datasetView.do',
'[4] 기상청. 종관기상관측 ASOS 시간자료. 서울 108 지점 2025년. https://data.kma.go.kr/data/grnd/selectAsosList.do?pgmNo=36',
'[5] Microsoft. LightGBM 공식 구현 및 파라미터 문서. https://github.com/microsoft/LightGBM ; CatBoost 공식 구현. https://github.com/catboost/catboost',
'[6] scikit-learn 개발진. Probability calibration 및 Tuning the decision threshold for class prediction. https://scikit-learn.org/stable/modules/calibration.html ; https://scikit-learn.org/stable/modules/classification_threshold.html',
'[7] Li 등. Diffusion Convolutional Recurrent Neural Network Data-Driven Traffic Forecasting. ICLR 2018. 공식 구현 https://github.com/liyaguang/DCRNN . 주변 흐름 활용의 조사 근거이며 본 과제에서 원모델 전체 재현은 하지 않음.',
'[8] Zeng 등. Are Transformers Effective for Time Series Forecasting. AAAI 2023. 공식 구현 https://github.com/cure-lab/LTSF-Linear . 시간 특징 설계의 조사 근거이며 본 과제의 원논문 재현 실적과 구분.',
'[9] P_Project 연구 산출물. config/improvement_v3.json ; artifacts/improvement/final/results.json ; selection_frozen_before_test.json ; reports/improvement_results.md ; reports/method_search/research_design.md . 방법별 원문·확인 이력·해시는 같은 폴더 sources.json 및 primary_sources.json.',
'[10] 가천대학교 인공지능학과. 2026 AI P 실무프로젝트 평가 260921. 2쪽 평가 항목과 비율, 3쪽 제출 및 발표 요건, 4쪽 교수 멘토링 회의록.',
'[11] 가천대학교 인공지능학과. 2026 AI P 실무프로젝트 OT. 3쪽 수업 목표, 8~12쪽 운영 및 멘토링, 15쪽 일정, 17쪽 GitHub 및 Wiki 운영.',
'[12] AI P 실무프로젝트 과제 계획서 양식. 기본정보·수행기간·본문 2~7항·기업 멘토링 기재 및 증빙 요건.',
]
for text in refs:p(text,10)
h('자료를 확인하는 순서',2)
p('핵심 수치는 [9]의 최종 JSON을 기준으로 하고, 이전 M1 결과는 artifacts/evaluation/model_results.json과 recommendation_results.json에서 별도로 확인한다. 데이터 검수는 artifacts/evaluation/data_quality.json과 artifacts/improvement/source_audit.json 및 weather_audit.json에 보관한다. 화면·API·실행 명령은 README와 docs/demo_v2_guide.md를 따른다.',10)
p('GitHub: https://github.com/Honggyu-Chun/P_Project\n인수 기준: transfer/local-codex-20261020, 커밋 81304e2f98510c14ef995ce71562716f5db43966',9.5)

# Serialize only document.xml; all other package parts remain byte-identical.
buf=BytesIO();doc.save(buf)
with zipfile.ZipFile(buf) as built: new_document=built.read('word/document.xml')
with zipfile.ZipFile(BytesIO(source)) as original,zipfile.ZipFile(OUT,'w',zipfile.ZIP_DEFLATED) as final:
    for entry in original.infolist():final.writestr(entry,new_document if entry.filename=='word/document.xml' else original.read(entry.filename))
with zipfile.ZipFile(OUT) as final:
    assert all(hashlib.sha256(final.read(n)).hexdigest()==v['sha256'] for n,v in inventory.items() if v['policy']=='preserve')
assert REF.read_bytes()==source
assert etree.tostring(doc.sections[0]._sectPr,method='c14n',exclusive=True)==etree.tostring(section,method='c14n',exclusive=True)
text='\n'.join(a.text for a in doc.paragraphs)
for t in doc.tables:
    for r in t.rows:text+='\n'+' | '.join(c.text for c in r.cells)
(TMP/'proposal_text.txt').write_text(text,encoding='utf-8')
print(json.dumps({'output':str(OUT),'sha256':hashlib.sha256(OUT.read_bytes()).hexdigest(),'source_unchanged':True,'preserve_only_parts_unchanged':True,'paragraphs':len(doc.paragraphs),'tables':len(doc.tables)},ensure_ascii=True,indent=2))
