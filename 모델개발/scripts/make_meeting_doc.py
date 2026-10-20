"""Create the meeting handout from the project's already verified research records."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables'
OUT.mkdir(exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.5)
sec.left_margin, sec.right_margin = Cm(1.7), Cm(1.7)
sec.header_distance, sec.footer_distance = Cm(.7), Cm(.7)
FONT = 'Noto Sans CJK KR'
NAVY, TEAL, GRAY = '16324B', '007E87', '536375'

def font(run, size=None, bold=False, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    if size: run.font.size = Pt(size)
    run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)
    return run

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    st = doc.styles[name]
    st.font.name = FONT
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
normal = doc.styles['Normal']
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.16
for name, size in [('Title', 23), ('Heading 1', 17), ('Heading 2', 11.5)]:
    st = doc.styles[name]
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(NAVY)
    st.paragraph_format.space_before = Pt(9 if name != 'Title' else 0)
    st.paragraph_format.space_after = Pt(6)

header = sec.header.paragraphs[0]
font(header.add_run('AI P-실무프로젝트  /  주제 검토 회의'), 8, color=GRAY)
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer = sec.footer.paragraphs[0]
font(footer.add_run('2026.10.19  ·  데이터 확인 결과와 실험 계획을 구분한 회의자료'), 8, color=GRAY)
font(footer.add_run('     |     '), 8, color=GRAY)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)

def p(text, bold=False, size=None, color=None):
    para = doc.add_paragraph()
    font(para.add_run(text), size, bold, color)
    return para

def h(text): doc.add_heading(text, 2)

def page(kicker, title):
    new_page = len(doc.paragraphs) > 1
    para = p(kicker, True, 9, TEAL)
    para.paragraph_format.page_break_before = new_page
    para.paragraph_format.keep_with_next = True
    doc.add_heading(title, 0 if kicker.startswith('01') else 1)

def callout(label, text):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'EAF5F5'); c._tc.get_or_add_tcPr().append(shade)
    q = c.paragraphs[0]
    q.paragraph_format.space_after = Pt(6); q.paragraph_format.space_before = Pt(5)
    font(q.add_run(label + '  '), 10, True, TEAL)
    font(q.add_run(text), 10)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def table(headers, rows, widths=None, size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths:
        for c, w in zip(t.columns, widths): c.width = Cm(w)
    for i, text in enumerate(headers):
        c = t.rows[0].cells[i]
        if widths: c.width = Cm(widths[i])
        font(c.paragraphs[0].add_run(text), size, True, 'FFFFFF')
        sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), NAVY); c._tc.get_or_add_tcPr().append(sh)
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for j, row in enumerate(rows):
        cells = t.add_row().cells
        for i, text in enumerate(row):
            c = cells[i]
            if widths: c.width = Cm(widths[i])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            q = c.paragraphs[0]
            q.paragraph_format.space_before = Pt(3)
            q.paragraph_format.space_after = Pt(4)
            q.paragraph_format.line_spacing = 1.12
            font(q.add_run(text), size)
            if j % 2 == 0:
                sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'F1F5F8'); c._tc.get_or_add_tcPr().append(sh)
        cant = OxmlElement('w:cantSplit'); t.rows[-1]._tr.get_or_add_trPr().append(cant)
    return t

def link(para, label, url):
    part = para.part
    rid = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hl = OxmlElement('w:hyperlink'); hl.set(qn('r:id'), rid)
    r = OxmlElement('w:r'); rp = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts'); rf.set(qn('w:ascii'), FONT); rf.set(qn('w:eastAsia'), FONT); rp.append(rf)
    col = OxmlElement('w:color'); col.set(qn('w:val'), TEAL); rp.append(col)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '17'); rp.append(sz)
    r.append(rp); txt = OxmlElement('w:t'); txt.text = label; r.append(txt); hl.append(r); para._p.append(hl)

page('01  /  회의의 출발점', '서울 전역 따릉이\n대여량 예측, 진행할 수 있을까?')
p('2026년 10월 19일  |  주제 선정·서비스 범위 협의용', size=9, color=GRAY)
callout('현재 판단', '핵심 데이터가 실제로 존재하고 서로 연결되어, 프로젝트를 시작할 수 있다. 예측 정확도 향상은 아직 검증하지 않았다.')
h('제안하는 주제와 핵심 질문')
p('주제: 서울 대여소별 시간대별 이용량 예측 및 운영 지원 대시보드', True)
p('“이 대여소의 과거 이용 패턴에 주변·유사 대여소의 패턴을 더하면, 어느 지역·시간대에서 예측 오차가 줄어드는가?”를 확인한다. 공연·행사 정보는 필수로 넣지 않는다.')
h('실제 파일을 열어 확인한 핵심 데이터')
table(['자료·대상 기간', '직접 확인한 결과', '프로젝트에서 쓰는 방법'], [
    ('따릉이 시간대별 이용정보 [1]\n2025년 1월 전체 CSV', '1,610,875행 / 대여소 ID 2,730개\n이용건수 합계 1,704,190건\n서울 전체 기준 744개 시간대 확인', '날짜×시간×대여소별 정답 생성.\n성별·연령·이용권으로 나뉜 행의 “이용건수”를 합산한다.'),
    ('대여소 위치 정보 [2]\n2024년 12월 XLSX', '이용 ID 2,730개 중 2,728개 연결\n그중 2,674곳에 500m 이웃 존재\n500m 이웃 수 중앙값 5곳', '위도·경도로 지도와 주변 대여소 특성 생성. 미연결 ID 4721·9980은 별도 확인한다.'),
    ('기상청 ASOS 서울 108 [3]\n2025년 연간 CSV', '8,760행 / 27개 항목\n1월 날씨와 이용정보 744/744시간 연결\n1월 기온·풍속·습도 결측 0', '기온·습도·풍속 등을 후보 입력으로 사용. 예측 시점에 실제 이용할 수 있는 날씨만 넣는다.'),
], [4.1, 6.5, 7.0], 9)
h('“어떻게 확인했나?”를 한 문장으로')
p('공식 사이트에서 원본 확보 → 압축·인코딩·열 구조 확인 → 날짜/ID/결측/건수 집계 → 대여소 ID와 위경도 연결 → 이용 시각과 기상 시각 연결. 단순히 목록만 검색한 결과가 아니다.')
p('주의: 744시간은 서울 전체에 기록이 있는 시간의 수다. 모든 대여소에 744시간 기록이 있다는 뜻은 아니다. 위 대여소 수는 2025년 자료 기준이며 현재 운영 개수와 다르다.', size=9, color=GRAY)

page('02  /  데이터 선택과 서비스', '핵심 3종부터 시작하고, 보조 자료는 선택한다')
h('추가로 검색하고 실제 내용을 확인한 자료')
table(['자료', '실제 확인 결과', '채택 판단'], [
    ('지하철 일별 승하차 [4]', '2025년 1월 CSV 19,129행. 날짜·노선·역·일별 승하차 수. 원본의 빈 마지막 필드를 확인 후 정리했다.', '보조 후보. 안내상 D+3 갱신이므로 같은 날의 확정 승객 수를 미리 넣을 수 없다. 역 위치 연결은 추가 작업.'),
    ('지하철 시간대별 승하차 [5]', '2025-01 JSON 621행·52필드. 시간대별 값은 “월간 누계”이며 일별 시간 관측이 아니다.', '후순위. 이미 공개가 끝난 이전 월 패턴만 사용. 시험 월 전체 합계를 넣으면 미래 정보가 섞인다.'),
    ('서울 주요 공원현황 [6]', '2026 하반기 XLSX: 155개 공원, 20항목. 넓은 서울 주변 범위 내 유효 좌표 141개.', '거리 특성 후보. 2025년 당시 상태를 보장하지 않고 모든 소공원을 포함하지 않아 우선 제외 가능.'),
    ('UCI Seoul Bike [7]', '실제 ZIP: 8,760행·14항목. 2017.12~2018.11. 대여소 ID 없음. CC BY 4.0 확인.', '연습·보조 벤치마크용. 대여소별 지도 서비스의 주데이터로 대체할 수 없다.'),
    ('공휴일·임시공휴일 [8]', '공식 API 문서 확인. 키 없는 실제 호출은 오류 응답으로 정상 목록 미확보. 오픈소스 달력은 별도 확인.', '최신 공식 2025년 결과 파일만 보완 필요. 오픈소스 파생 달력을 공식 원본으로 표시하지 않는다.'),
], [3.8, 7.0, 6.8], 8.8)
h('서비스 구상: 운영 담당자가 보는 시간대별 전망 지도')
p('화면 ① 서울 지도에서 날짜·시간을 선택하고 예상 대여량이 많은 대여소를 확인한다.\n화면 ② 대여소를 누르면 시간대별 예상 대여 건수와 과거 이용 패턴을 본다.\n화면 ③ 과거 날짜에서는 실제값·기본 모델·개선 모델의 차이와 오차를 비교한다.')
p('구현 흐름: 공개 파일 → 대여소×시간 집계 → 모델 → 예측값 일괄 저장 → API → 지도·그래프. Python·LightGBM, 간단한 API와 Streamlit 화면으로 시작할 수 있다.', size=9)
h('처음부터 지킬 서비스 범위')
p('공개 대여이력은 반기 갱신이므로, 첫 제품은 과거 자료를 활용한 주기적 전망·과거 재생 시연으로 정한다. 현재 직전 1시간 이용량을 받는 실시간 서비스는 별도 공급 경로가 필요하다.')
p('예측 대상은 실제 성사된 “대여 건수”다. 재고 수량, 자전거 부족으로 이용하지 못한 잠재 수요, 재배치 효과까지 측정한 것은 아니다.', True, 9)
p('신뢰 근거: 핵심 자료는 서울시·기상청 공식 배포본이다. 원본 출처·다운로드 기록·해시와 검증 결과 JSON을 저장했다. 서울시 자료의 출처표시 조건을 확인했으며 제공 기관과 자료 버전을 명시한다.', size=8.5, color=GRAY)

page('03  /  앞으로 검증할 것', '정확도가 좋아지는지는 공정한 비교로 확인한다')
callout('완료와 계획 구분', '이번에 끝난 것은 데이터 확보·내용 검사·시간/위치 연결이다. 아래 모델 학습·성능 비교와 서비스 구현은 다음 단계 계획이다.')
h('동일한 정답·기간을 쓰는 3단계 비교')
table(['구성', '입력과 역할'], [
    ('B0 · 단순 기준', '같은 대여소의 사용 가능한 과거 동일 요일·시간 평균 또는 전주 값. 실제 공개 시점에 맞는 과거만 사용한다.'),
    ('B1 · 주 비교 기준', 'LightGBM + 해당 대여소의 과거 패턴 + 시간·요일·휴일 + 이용 가능한 날씨. 날씨·휴일만 추가했다고 차별성으로 주장하지 않는다.'),
    ('M1 · 개선 후보', 'B1과 같은 LightGBM에 주변 대여소 거리·과거 패턴 및 유사 대여소 그룹 특성을 추가한다. “이웃 정보가 도움이 되는가”를 분리해서 검증한다.'),
], [4, 13.6], 9)
h('비교 규칙과 판단 기준')
p('① 여러 달의 자료를 추가 확보한 뒤, 과거 → 미래 순서로 학습·검증·시험 기간을 고정한다. 시험 정답을 보고 기간이나 특성을 바꾸지 않는다. 현재 한 달 확인만으로 연중 성능을 주장하지 않는다.')
p('② 같은 시험 행과 기간을 사용하고 B1/M1은 같은 모델·튜닝 예산으로 비교한다. 평균·군집·결측 대체 규칙도 학습 자료에서만 만든다. 주변의 현재 1시간 이용량 대신 공개된 과거 패턴부터 쓴다.')
p('③ 주지표 MAE(시간당 평균 몇 건 틀렸는지), 보조지표 RMSE(큰 오차에 더 민감)를 보고한다. 전체뿐 아니라 출퇴근·고수요 시간, 대여소별 오차와 표본 수도 비교한다.')
h('학습 전에 해결할 두 가지')
p('• 기록 없음 ≠ 확인된 0건: 운영 중 0건인지, 자료 누락·미운영인지 구분하는 규칙과 평가 대상 표시가 필요하다. 대여가 있었던 시간만 평가하거나 모든 빈칸을 바로 0으로 채우지 않는다.\n• 관측 날씨 ≠ 미래 예보: 받은 연간 파일은 사후 관측본이다. 미래의 실제 날씨를 입력하면 안 된다. 강수 공란 690/744칸도 관측 규칙 확인 전 0으로 처리하지 않는다.', size=9)
h('과제 평가 요소에 연결할 증거')
table(['공식 평가 안내 기준', '준비할 결과물'], [
    ('사전 40% · 적합성·창의성\n계획서·발표·Q&A', '실제 데이터와 해결할 문제를 제시. 흔한 예측 주제라는 약점은 “주변 정보가 언제·어디서 도움이 되는가”와 운영 화면으로 보완한다.'),
    ('최종 60% · 성능·Baseline\n구현·적용·문서·발표·동료평가', '동일 조건 비교표, 오차 분석, 작동하는 지도/그래프, 출처·재현 방법·GitHub/Wiki, 4명 담당 산출물과 기여 기록을 남긴다.'),
], [5.6, 12.0], 8.8)

page('04  /  오늘 결정할 내용', '범위를 정하고, 사전평가 준비를 시작한다')
h('회의 안건 · 20~30분')
table(['결정할 항목', '권장안 / 회의 메모'], [
    ('1. 주제 채택 · 5분', '□ 채택  □ 수정  □ 보류\n서울 대여소별 이용량 예측 + 주변 패턴의 효과 검증'),
    ('2. 첫 서비스 범위 · 5분', '핵심 3종(이용량·위치·날씨) + 달력부터. 지도·예측 그래프·과거 비교 화면.\n추가 데이터 선택: □ 지하철  □ 공원  □ 우선 제외'),
    ('3. 실험·역할 · 5~10분', 'B0/B1/M1 비교, 0건 판정·입력 공개 시점 규칙부터 합의.\n담당: 모델 ______ / 데이터·API ______ / 데이터 QA ______ / 화면·문서 ______'),
    ('4. 일정·보완 · 5분', '사전평가까지 확정할 주제·계획서·발표 담당 및 공휴일 자료 확보 담당 결정.\n추가 의견: __________________________________________________'),
], [4.1, 13.5], 9)
h('가까운 제출 일정')
p('10/19~10/21: 주제·서비스 범위·담당 확정, 0건/지연 처리 원칙과 추가 월 확보 계획 정리.\n10/22~10/24: 계획서·발표자료 작성, 데이터 근거와 화면 초안·비교 실험 설계 점검.\n10/25(일) 23:59: 사전평가 제출 — PPT·PDF·10분 이내 영상 URL TXT·계획서 WORD를 ZIP으로 팀장 1명이 제출.\n12/13(일) 23:59: 최종 제출. 중간에는 학습·비교·서비스 통합과 멘토 피드백을 반영한다.', size=9)
h('현재 추가로 받으면 되는 자료는 공식 공휴일 목록 하나')
p('공공데이터포털 「한국천문연구원_특일 정보」[8] → 필요 시 활용신청 → “공휴일 정보 조회”. solYear=2025, numOfRows=100, pageNo=1, solMonth는 비워 연간 조회. XML/JSON 응답 또는 locdate·dateName·isHoliday가 있는 CSV를 받으면 된다. 인증키가 아닌 결과 파일만 공유한다.', size=8.7)
p('2025년 1/27·6/3 추가 지정 여부까지 확인한다. 이전 공개 달력 표는 최신 공식 목록으로 쓰기 어려웠다. 따릉이·위치·기상·지하철·공원 파일은 이미 확보해 다시 받을 필요가 없다.', size=8.7)
h('공식 출처 · 제목을 누르면 원문으로 이동')
refs = [
    ('[1] 서울시 따릉이 시간대별 이용정보 (OA-15245)', 'https://data.seoul.go.kr/dataList/OA-15245/F/1/datasetView.do'),
    ('[2] 서울시 공공자전거 대여소 정보 (OA-13252)', 'https://data.seoul.go.kr/dataList/OA-13252/F/1/datasetView.do'),
    ('[3] 기상청 종관기상관측 ASOS · 서울 108', 'https://data.kma.go.kr/data/grnd/selectAsosList.do?pgmNo=36'),
    ('[4] 서울시 지하철 일별 승하차 (OA-12914)', 'https://data.seoul.go.kr/dataList/OA-12914/F/1/datasetView.do'),
    ('[5] 서울시 지하철 시간대별 승하차 (OA-12252)', 'https://data.seoul.go.kr/dataList/OA-12252/S/1/datasetView.do'),
    ('[6] 서울시 주요 공원현황 (OA-394)', 'https://data.seoul.go.kr/dataList/OA-394/S/1/datasetView.do'),
    ('[7] UCI Seoul Bike Sharing Demand', 'https://archive.ics.uci.edu/dataset/560/seoul+bike+sharing+demand'),
    ('[8] 한국천문연구원 특일 정보 · 공공데이터포털', 'https://www.data.go.kr/data/15012690/openapi.do'),
]
for i in range(0, len(refs), 2):
    q = doc.add_paragraph(); q.paragraph_format.space_after = Pt(2)
    link(q, *refs[i]); font(q.add_run('   /   '), 8, color=GRAY); link(q, *refs[i+1])
p('확인 기록: reports/citywide_feasibility_brief.md 및 공간·기상·지하철 검증 JSON, 다운로드 명세. 기존 소규모 행사 실험은 이번 서울 전역 모델의 성능 근거로 사용하지 않았다.', size=8, color=GRAY)
p('평가·일정 근거: 첨부 「2026_AI P-실무프로젝트_평가_260921.pdf」 2~3쪽. 신규 딥러닝 알고리즘 필수 조건은 확인되지 않았다. 성능 개선율·품절 감소 효과는 아직 검증하지 않았다.', size=8, color=GRAY)

doc.core_properties.title = '서울 전역 따릉이 대여량 예측 — 주제 검토 회의자료'
doc.core_properties.subject = '실제 데이터 검증 결과·서비스 구상·평가 설계·회의 안건'
doc.core_properties.author = 'P_Project'
doc.core_properties.keywords = '따릉이, 수요 예측, 회의자료, 데이터 검증, AI P-실무프로젝트'
path = OUT / '서울_따릉이_수요예측_주제검토_회의자료_20261019.docx'
doc.save(path)
print(path)
